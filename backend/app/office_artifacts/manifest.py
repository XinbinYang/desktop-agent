from __future__ import annotations

import hashlib
import json
import re
import shutil
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string

from app.office_artifacts.ppt_extract import extract_presentation
from app.office_artifacts.renderers import render_office_previews
from app.office_artifacts.xlsx_extract import extract_workbook, scan_formula_errors
from app.runtime_paths import runtime_dir


MANIFEST_VERSION = 2
OFFICE_EXTENSIONS = {
    ".xlsx": "excel",
    ".xls": "excel",
    ".pptx": "ppt",
    ".ppt": "ppt",
}


class OfficeArtifactNotFound(FileNotFoundError):
    pass


class OfficeArtifactEditError(ValueError):
    pass


def _safe_segment(value: str, fallback: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value or "").strip()).strip("._")
    return cleaned[:120] or fallback


def _is_relative_to(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _preview_url(path: Path) -> str:
    preview_root = runtime_dir("preview").resolve()
    rel = path.resolve().relative_to(preview_root)
    return "/preview/" + "/".join(quote(part) for part in rel.parts)


def _artifact_root() -> Path:
    path = runtime_dir("office_artifacts")
    path.mkdir(parents=True, exist_ok=True)
    return path


def _session_root(session_id: str) -> Path:
    path = _artifact_root() / _safe_segment(session_id or "default", "default")
    path.mkdir(parents=True, exist_ok=True)
    return path


def _manifest_path(artifact_id: str, session_id: str | None = None) -> Path:
    safe_id = _safe_segment(artifact_id, "")
    if not safe_id or safe_id != artifact_id:
        raise OfficeArtifactNotFound(f"Office artifact was not found: {artifact_id}")
    if session_id:
        return _session_root(session_id) / safe_id / "manifest.json"
    root = _artifact_root()
    matches = list(root.glob(f"*/{safe_id}/manifest.json"))
    if not matches:
        raise OfficeArtifactNotFound(f"Office artifact was not found: {artifact_id}")
    return matches[0]


def office_kind_for_path(path: str | Path, explicit: str = "") -> str:
    kind = str(explicit or "").strip().lower()
    if kind in {"excel", "xlsx", "xls", "spreadsheet", "workbook"}:
        return "excel"
    if kind in {"ppt", "pptx", "powerpoint", "presentation", "deck"}:
        return "ppt"
    return OFFICE_EXTENSIONS.get(Path(path).suffix.lower(), "office")


def _mime_for_kind(kind: str, path: Path) -> str:
    if path.suffix.lower() == ".xlsx":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if path.suffix.lower() == ".xls":
        return "application/vnd.ms-excel"
    if path.suffix.lower() == ".pptx":
        return "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    if path.suffix.lower() == ".ppt":
        return "application/vnd.ms-powerpoint"
    return "application/octet-stream"


def _available_actions(kind: str) -> list[str]:
    base = ["open_native", "reveal_file", "search", "zoom"]
    if kind == "excel":
        return base + ["switch_sheets", "formula_bar", "select_copy", "basic_filter", "simple_value_edit"]
    if kind == "ppt":
        return base + ["thumbnail_nav", "page_jump", "fit_width", "speaker_notes", "text_summary"]
    return base


def _copy_to_preview(source: Path, *, session_id: str, artifact_id: str, title: str) -> Path:
    preview_root = runtime_dir("preview").resolve()
    source = source.resolve()
    if _is_relative_to(source, preview_root):
        return source

    target_dir = preview_root / _safe_segment(session_id or "default", "default") / "office" / _safe_segment(artifact_id, "artifact")
    target_dir.mkdir(parents=True, exist_ok=True)
    display_name = _safe_segment(title or source.name, source.name)
    if not display_name.lower().endswith(source.suffix.lower()):
        display_name = f"{display_name}{source.suffix}"
    target = target_dir / display_name
    if source != target.resolve():
        shutil.copy2(source, target)
    return target.resolve()


def _qa_for_workbook(workbook_payload: dict[str, Any], render_issues: list[str]) -> dict[str, Any]:
    formula_errors = scan_formula_errors(workbook_payload)
    issues: list[str] = []
    if formula_errors:
        issues.append(f"{len(formula_errors)} formula error(s) detected.")
    if not workbook_payload.get("sheet_count"):
        issues.append("Workbook has no sheets.")
    warnings: list[str] = []
    if not workbook_payload.get("table_count"):
        warnings.append("No structured Excel table was detected.")
    if not workbook_payload.get("chart_count"):
        warnings.append("No native Excel chart was detected.")
    qa_status = "failed" if formula_errors or not workbook_payload.get("sheet_count") else "warnings" if warnings else "passed"
    return {
        "qa_status": qa_status,
        "sheet_count": workbook_payload.get("sheet_count", 0),
        "formula_count": workbook_payload.get("formula_count", 0),
        "table_count": workbook_payload.get("table_count", 0),
        "chart_count": workbook_payload.get("chart_count", 0),
        "validation_count": workbook_payload.get("validation_count", 0),
        "formula_errors": formula_errors,
        "issues": issues,
        "warnings": warnings,
        "render_issues": render_issues,
    }


def _qa_for_presentation(presentation_payload: dict[str, Any], render_issues: list[str], preview_count: int) -> dict[str, Any]:
    layout_errors = presentation_payload.get("layout_errors") or []
    layout_warnings = presentation_payload.get("layout_warnings") or []
    slide_count = int(presentation_payload.get("slide_count") or 0)
    issues = list(layout_errors)
    if slide_count <= 0:
        issues.append("Presentation has no slides.")
    if preview_count <= 0:
        issues.append("No slide preview could be rendered.")
    qa_status = "failed" if issues else "warnings" if layout_warnings else "passed"
    return {
        "qa_status": qa_status,
        "slide_count": slide_count,
        "layout_errors": len(layout_errors),
        "layout_warnings": len(layout_warnings),
        "issues": issues,
        "warnings": list(layout_warnings),
        "render_issues": render_issues,
    }


def _merge_qa(kind: str, generated: dict[str, Any], provided: dict[str, Any] | None) -> dict[str, Any]:
    if not provided:
        return {kind: generated}
    merged = dict(provided)
    existing = merged.get(kind)
    if isinstance(existing, dict):
        merged[kind] = {**generated, **existing}
    else:
        merged[kind] = generated
    return merged


def _write_manifest(path: Path, manifest: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def register_office_artifact(
    path: str | Path,
    *,
    artifact: dict[str, Any] | None = None,
    session_id: str = "default",
    source_tool: str = "",
    tool_call_id: str = "",
    kind: str = "",
    qa_summary: dict[str, Any] | None = None,
    render: bool = True,
) -> dict[str, Any]:
    source = Path(path).expanduser().resolve()
    if not source.exists() or not source.is_file():
        raise OfficeArtifactNotFound(f"Office artifact file does not exist: {source}")

    artifact_payload = dict(artifact or {})
    artifact_id = _safe_segment(str(artifact_payload.get("id") or f"artifact_{uuid.uuid4().hex[:12]}"), "artifact")
    title = str(artifact_payload.get("title") or source.name)
    session = session_id or "default"
    manifest_dir = _session_root(session) / artifact_id
    manifest_dir.mkdir(parents=True, exist_ok=True)
    published_path = _copy_to_preview(source, session_id=session, artifact_id=artifact_id, title=title)
    kind_value = office_kind_for_path(published_path, kind or str(artifact_payload.get("kind") or ""))

    workbook_payload: dict[str, Any] | None = None
    presentation_payload: dict[str, Any] | None = None
    render_result: dict[str, Any] = {"engine": "none", "issues": [], "previews": [], "preview_paths": []}
    generated_qa: dict[str, Any] = {"qa_status": "warnings", "issues": ["Unsupported Office artifact type."]}

    if kind_value == "excel":
        workbook_payload = extract_workbook(published_path)
        if render:
            render_result = render_office_previews(
                published_path,
                kind="excel",
                output_dir=runtime_dir("preview") / _safe_segment(session, "default") / "office" / artifact_id / "previews",
                workbook_payload=workbook_payload,
            )
        generated_qa = _qa_for_workbook(workbook_payload, list(render_result.get("issues") or []))
    elif kind_value == "ppt":
        presentation_payload = extract_presentation(published_path)
        if render:
            render_result = render_office_previews(
                published_path,
                kind="ppt",
                output_dir=runtime_dir("preview") / _safe_segment(session, "default") / "office" / artifact_id / "previews",
                presentation_payload=presentation_payload,
            )
        generated_qa = _qa_for_presentation(
            presentation_payload,
            list(render_result.get("issues") or []),
            len(render_result.get("previews") or []),
        )

    file_url = artifact_payload.get("url") if isinstance(artifact_payload.get("url"), str) else _preview_url(published_path)
    size = int(published_path.stat().st_size)
    mime_type = str(artifact_payload.get("mime_type") or artifact_payload.get("mimeType") or _mime_for_kind(kind_value, published_path))
    previews = list(render_result.get("previews") or [])
    manifest = {
        "version": MANIFEST_VERSION,
        "artifact_id": artifact_id,
        "kind": kind_value,
        "title": title,
        "session_id": session,
        "created_at": _now_iso(),
        "source_tool": source_tool or str(artifact_payload.get("source") or artifact_payload.get("source_tool") or ""),
        "tool_call_id": tool_call_id or str(artifact_payload.get("tool_call_id") or ""),
        "file_path": str(published_path),
        "file_url": file_url,
        "mime_type": mime_type,
        "size": size,
        "sha256": _sha256(published_path),
        "files": [{
            "role": kind_value,
            "name": published_path.name,
            "path": str(published_path),
            "url": file_url,
            "mime_type": mime_type,
            "size": size,
        }],
        "workbook": workbook_payload,
        "presentation": presentation_payload,
        "previews": previews,
        "render_engine": render_result.get("engine") or "none",
        "render_issues": list(render_result.get("issues") or []),
        "qa_summary": _merge_qa(kind_value, generated_qa, qa_summary),
        "available_actions": _available_actions(kind_value),
    }
    manifest = {key: value for key, value in manifest.items() if value is not None}
    manifest_path = manifest_dir / "manifest.json"
    _write_manifest(manifest_path, manifest)

    enriched = dict(artifact_payload)
    enriched.update({
        "id": artifact_id,
        "type": "office",
        "title": title,
        "url": file_url,
        "path": str(published_path),
        "kind": kind_value,
        "mime_type": mime_type,
        "size": size,
        "sha256": manifest["sha256"],
        "office_manifest_id": artifact_id,
        "manifest_path": str(manifest_path),
        "manifest_url": f"/api/office/{artifact_id}/manifest",
        "viewer_manifest_url": f"/api/office/{artifact_id}/manifest",
        "previews": previews,
        "qa_summary": manifest["qa_summary"],
        "engine": manifest["render_engine"],
        "render_issues": manifest["render_issues"],
        "available_actions": manifest["available_actions"],
    })
    if workbook_payload:
        enriched["workbook"] = workbook_payload
    if presentation_payload:
        enriched["presentation"] = presentation_payload
    return enriched


def get_manifest(artifact_id: str, *, session_id: str | None = None) -> dict[str, Any]:
    path = _manifest_path(artifact_id, session_id=session_id)
    if not path.exists() or not path.is_file():
        raise OfficeArtifactNotFound(f"Office artifact was not found: {artifact_id}")
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise OfficeArtifactNotFound(f"Office artifact manifest is unavailable: {artifact_id}") from exc
    if int(data.get("version") or 0) < MANIFEST_VERSION:
        raise OfficeArtifactNotFound(f"Unsupported Office artifact manifest version: {artifact_id}")
    return data


def update_xlsx_simple_edits(
    artifact_id: str,
    edits: list[dict[str, Any]],
    *,
    output_name: str = "",
    session_id: str | None = None,
) -> dict[str, Any]:
    manifest = get_manifest(artifact_id, session_id=session_id)
    if manifest.get("kind") != "excel":
        raise OfficeArtifactEditError("Simple edits are only supported for Excel artifacts.")
    if not edits:
        raise OfficeArtifactEditError("At least one cell edit is required.")

    source = Path(str(manifest.get("file_path") or "")).expanduser().resolve()
    preview_root = runtime_dir("preview").resolve()
    if not _is_relative_to(source, preview_root):
        raise OfficeArtifactEditError("Workbook path is outside preview storage.")

    workbook = load_workbook(source, data_only=False)
    for edit in edits:
        sheet_name = str(edit.get("sheet") or "").strip()
        address = str(edit.get("cell") or edit.get("address") or "").strip().upper()
        if not sheet_name or sheet_name not in workbook.sheetnames:
            raise OfficeArtifactEditError(f"Unknown sheet: {sheet_name or '(blank)'}")
        try:
            coordinate_from_string(address)
        except Exception as exc:
            raise OfficeArtifactEditError(f"Invalid cell address: {address}") from exc
        cell = workbook[sheet_name][address]
        if isinstance(cell.value, str) and cell.value.startswith("="):
            raise OfficeArtifactEditError(f"{sheet_name}!{address} contains a formula and cannot be overwritten by simple edits.")
        value = edit.get("value")
        if isinstance(value, str) and value.startswith("="):
            raise OfficeArtifactEditError("Simple edits only accept literal values, not formulas.")
        cell.value = value

    original_name = Path(output_name or manifest.get("title") or source.name).name
    if not original_name.lower().endswith(".xlsx"):
        original_name += ".xlsx"
    new_id = f"artifact_{uuid.uuid4().hex[:12]}"
    session = str(manifest.get("session_id") or session_id or "default")
    out_dir = runtime_dir("preview") / _safe_segment(session, "default") / "office" / new_id
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / _safe_segment(original_name, "edited.xlsx")
    workbook.save(out_path)

    artifact = {
        "id": new_id,
        "type": "office",
        "title": out_path.name,
        "url": _preview_url(out_path),
        "path": str(out_path),
        "kind": "excel",
        "source": "office_xlsx_simple_edits",
        "size": out_path.stat().st_size,
    }
    return register_office_artifact(
        out_path,
        artifact=artifact,
        session_id=session,
        source_tool="office_xlsx_simple_edits",
        kind="excel",
    )
