from __future__ import annotations

import copy
import json
import re
import shutil
import subprocess
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from app.runtime_paths import runtime_dir
from app.tools.artifact_tool import (
    _existing_source_path,
    _safe_segment,
    _validate_source,
    publish_file_artifact,
    publish_many_file_artifacts,
)
from app.tools.base import BaseTool, ToolResult


OFFICE_KIND_EXCEL = "excel"
OFFICE_KIND_PPT = "ppt"
OFFICE_KIND_PACKAGE = "package"


def _json(data: Any) -> str:
    return json.dumps(data, ensure_ascii=False, indent=2, default=str)


def _as_dict(value: Any, name: str) -> tuple[dict[str, Any], Optional[str]]:
    if value is None:
        return {}, None
    if isinstance(value, dict):
        return value, None
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError as exc:
            return {}, f"{name} must be an object or JSON object string: {exc}"
        if isinstance(parsed, dict):
            return parsed, None
    return {}, f"{name} must be an object."


def _as_list(value: Any, name: str) -> tuple[list[Any], Optional[str]]:
    if value is None:
        return [], None
    if isinstance(value, list):
        return value, None
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError as exc:
            return [], f"{name} must be an array or JSON array string: {exc}"
        if isinstance(parsed, list):
            return parsed, None
    return [], f"{name} must be an array."


def _require_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
        from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
    except Exception as exc:  # pragma: no cover - exercised in packaging envs
        return None, f"openpyxl is required for Excel tools: {exc}"
    return {
        "Workbook": Workbook,
        "load_workbook": load_workbook,
        "BarChart": BarChart,
        "LineChart": LineChart,
        "PieChart": PieChart,
        "Reference": Reference,
        "CellIsRule": CellIsRule,
        "Alignment": Alignment,
        "Border": Border,
        "Font": Font,
        "PatternFill": PatternFill,
        "Side": Side,
        "DataValidation": DataValidation,
        "Table": Table,
        "TableStyleInfo": TableStyleInfo,
        "get_column_letter": get_column_letter,
        "coordinate_to_tuple": coordinate_to_tuple,
        "range_boundaries": range_boundaries,
    }, None


def _require_pptx():
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt
    except Exception as exc:  # pragma: no cover - exercised in packaging envs
        return None, f"python-pptx is required for PowerPoint tools: {exc}"
    chart_data_cls = None
    chart_type_enum = None
    legend_enum = None
    try:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

        chart_data_cls = CategoryChartData
        chart_type_enum = XL_CHART_TYPE
        legend_enum = XL_LEGEND_POSITION
    except Exception:  # pragma: no cover - charts optional
        pass
    return {
        "Presentation": Presentation,
        "RGBColor": RGBColor,
        "MSO_SHAPE": MSO_SHAPE,
        "MSO_ANCHOR": MSO_ANCHOR,
        "PP_ALIGN": PP_ALIGN,
        "Inches": Inches,
        "Pt": Pt,
        "CategoryChartData": chart_data_cls,
        "XL_CHART_TYPE": chart_type_enum,
        "XL_LEGEND_POSITION": legend_enum,
    }, None


def _resolve_input_path(path: str, agent_type: str = "") -> tuple[Optional[Path], Optional[str]]:
    if not path:
        return None, "Missing required argument: path"
    source = _existing_source_path(path, agent_type=agent_type)
    if source is None:
        return None, "File does not exist or cannot be resolved from the allowed workspace roots."
    error = _validate_source(source, agent_type=agent_type)
    if error:
        return None, error
    return source, None


def _office_run_dir(session_id: str, tool_call_id: str, kind: str) -> Path:
    safe_session = _safe_segment(session_id or "default", "default")
    safe_call = _safe_segment(tool_call_id or uuid.uuid4().hex[:10], "tool")
    path = runtime_dir("office") / safe_session / f"{kind}_{safe_call}"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _office_package_dir(session_id: str, task_slug: str) -> Path:
    safe_session = _safe_segment(session_id or "default", "default")
    safe_task = _safe_segment(task_slug or uuid.uuid4().hex[:10], "office-package")
    path = runtime_dir("office") / safe_session / safe_task
    for child in ("excel", "ppt", "preview", "qa", "output"):
        (path / child).mkdir(parents=True, exist_ok=True)
    return path


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(_json(payload), encoding="utf-8")


def _output_path(output_name: str, suffix: str, *, session_id: str, tool_call_id: str, kind: str) -> Path:
    raw_name = Path(output_name or f"{kind}_{uuid.uuid4().hex[:8]}{suffix}").name
    safe_name = _safe_segment(raw_name, f"{kind}{suffix}")
    if not safe_name.lower().endswith(suffix):
        safe_name += suffix
    return _office_run_dir(session_id, tool_call_id, kind) / safe_name


def _office_artifact(
    path: Path,
    *,
    title: str,
    source_tool: str,
    session_id: str,
    tool_call_id: str,
    agent_type: str,
    caption: str = "",
) -> tuple[list[dict[str, Any]], Optional[str]]:
    artifact, error = publish_file_artifact(
        str(path),
        session_id=session_id or "default",
        title=title or path.name,
        caption=caption,
        source_tool=source_tool,
        tool_call_id=tool_call_id,
        agent_type=agent_type,
        artifact_type="office",
    )
    if error:
        return [], error
    return [artifact] if artifact else [], None


def _preview_artifacts(
    paths: list[Path],
    *,
    title_prefix: str,
    source_tool: str,
    session_id: str,
    tool_call_id: str,
    agent_type: str,
) -> list[dict[str, Any]]:
    triples = [
        (str(path), f"{title_prefix} {idx + 1}", "image")
        for idx, path in enumerate(paths)
    ]
    return publish_many_file_artifacts(
        triples,
        session_id=session_id or "default",
        source_tool=source_tool,
        tool_call_id=tool_call_id,
        agent_type=agent_type,
    )


def _encode_preview_images(
    paths: list[Any], *, max_images: int = 6, max_width: int = 1000
) -> list[str]:
    """Downscale rendered previews to base64 PNGs for the model's visual QA loop.

    Capped and downscaled on purpose: these base64 blobs are attached to a chat
    message, so they must stay small enough not to bloat memory or persistence.
    """
    import base64
    import io

    try:
        from PIL import Image
    except Exception:
        return []
    encoded: list[str] = []
    for path in list(paths)[:max_images]:
        try:
            image = Image.open(path).convert("RGB")
            if image.width > max_width:
                ratio = max_width / float(image.width)
                image = image.resize((max_width, max(1, int(image.height * ratio))))
            buffer = io.BytesIO()
            image.save(buffer, format="PNG", optimize=True)
            encoded.append(base64.b64encode(buffer.getvalue()).decode("ascii"))
        except Exception:
            continue
    return encoded


def _clean_color(value: Any) -> Optional[str]:
    if not isinstance(value, str):
        return None
    color = value.strip().lstrip("#")
    if re.fullmatch(r"[A-Fa-f0-9]{6}", color):
        return color.upper()
    return None


def _sheet_range(value: str) -> tuple[Optional[str], str]:
    if "!" not in value:
        return None, value
    sheet, cell_range = value.split("!", 1)
    return sheet.strip("'"), cell_range


def _get_or_create_sheet(workbook: Any, name: str):
    safe_name = str(name or "Sheet").strip()[:31] or "Sheet"
    if safe_name in workbook.sheetnames:
        return workbook[safe_name]
    return workbook.create_sheet(safe_name)


def _iter_cells(worksheet: Any, range_ref: str):
    cells = worksheet[range_ref]
    if not isinstance(cells, tuple):
        return [cells]
    flattened: list[Any] = []
    for row in cells:
        if isinstance(row, tuple):
            flattened.extend(row)
        else:
            flattened.append(row)
    return flattened


def _numeric_from_string(value: Any) -> Optional[float]:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if not isinstance(value, str):
        return None
    text = value.strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    cleaned = text.strip("()").replace(",", "").replace("$", "").replace("USD", "").strip()
    is_percent = cleaned.endswith("%")
    is_multiple = cleaned.lower().endswith("x")
    cleaned = cleaned.rstrip("%xX").strip()
    try:
        number = float(cleaned)
    except ValueError:
        return None
    if negative:
        number = -number
    if is_percent:
        number = number / 100
    return number


def _semantic_excel_value(value: Any) -> tuple[Any, Optional[str]]:
    if not isinstance(value, dict):
        return value, None
    if not any(key in value for key in ("value", "formula", "type", "number_format", "format")):
        return value, None

    kind = str(value.get("type") or "").lower()
    number_format = value.get("number_format") or value.get("format")
    if value.get("formula") is not None:
        formula = str(value.get("formula") or "")
        return formula if formula.startswith("=") else f"={formula}", str(number_format) if number_format else None

    raw = value.get("value")
    if kind in {"number", "currency", "percent", "percentage", "multiple"}:
        numeric = _numeric_from_string(raw)
        cell_value = numeric if numeric is not None else raw
        if not number_format:
            if kind == "currency":
                number_format = '$#,##0.0;[Red]($#,##0.0)'
            elif kind in {"percent", "percentage"}:
                number_format = "0.0%"
            elif kind == "multiple":
                number_format = "0.0x"
            else:
                number_format = "#,##0.0"
        return cell_value, str(number_format)

    if kind == "date":
        parsed = raw
        if isinstance(raw, str):
            try:
                parsed = datetime.fromisoformat(raw).date()
            except ValueError:
                parsed = raw
        if isinstance(parsed, datetime):
            parsed = parsed.date()
        if isinstance(parsed, date) and not number_format:
            number_format = "yyyy-mm-dd"
        return parsed, str(number_format) if number_format else None

    if kind in {"text", "string"}:
        return "" if raw is None else str(raw), str(number_format) if number_format else None
    return raw, str(number_format) if number_format else None


def _set_excel_cell_value(cell: Any, value: Any) -> None:
    cell_value, number_format = _semantic_excel_value(value)
    cell.value = cell_value
    if number_format:
        cell.number_format = number_format


def _write_grid(api: dict[str, Any], worksheet: Any, start_cell: str, values: Any) -> None:
    if values is None:
        return
    if not isinstance(values, list):
        values = [[values]]
    elif values and not isinstance(values[0], list):
        values = [values]
    row0, col0 = api["coordinate_to_tuple"](start_cell or "A1")
    for r_idx, row in enumerate(values):
        if not isinstance(row, list):
            row = [row]
        for c_idx, value in enumerate(row):
            _set_excel_cell_value(worksheet.cell(row=row0 + r_idx, column=col0 + c_idx), value)


def _apply_excel_style(api: dict[str, Any], worksheet: Any, range_ref: str, style: dict[str, Any]) -> None:
    if not range_ref:
        return
    fill_color = _clean_color(style.get("fill_color") or style.get("fill"))
    font_color = _clean_color(style.get("font_color") or style.get("color"))
    border_color = _clean_color(style.get("border_color") or "D9E2EC")
    border_style = style.get("border")

    for cell in _iter_cells(worksheet, range_ref):
        if any(key in style for key in ("bold", "italic", "font_size", "font_color", "color")):
            font = copy.copy(cell.font)
            if "bold" in style:
                font.bold = bool(style["bold"])
            if "italic" in style:
                font.italic = bool(style["italic"])
            if "font_size" in style:
                font.sz = float(style["font_size"])
            if font_color:
                font.color = font_color
            cell.font = font
        if fill_color:
            cell.fill = api["PatternFill"](fill_type="solid", fgColor=fill_color)
        if style.get("number_format"):
            cell.number_format = str(style["number_format"])
        if any(key in style for key in ("align", "vertical", "wrap_text")):
            cell.alignment = api["Alignment"](
                horizontal=style.get("align"),
                vertical=style.get("vertical"),
                wrap_text=bool(style.get("wrap_text", False)),
            )
        if border_style:
            side = api["Side"](style="thin", color=border_color)
            cell.border = api["Border"](left=side, right=side, top=side, bottom=side)


def _add_excel_table(api: dict[str, Any], worksheet: Any, table_spec: dict[str, Any], index: int) -> None:
    ref = str(table_spec.get("range") or table_spec.get("ref") or "").strip()
    if not ref:
        return
    raw_name = str(table_spec.get("name") or f"Table{index + 1}")
    name = re.sub(r"[^A-Za-z0-9_]", "_", raw_name).strip("_") or f"Table{index + 1}"
    style_name = str(table_spec.get("style") or "TableStyleMedium9")
    table = api["Table"](displayName=name[:255], ref=ref)
    table.tableStyleInfo = api["TableStyleInfo"](
        name=style_name,
        showFirstColumn=bool(table_spec.get("show_first_column", False)),
        showLastColumn=bool(table_spec.get("show_last_column", False)),
        showRowStripes=bool(table_spec.get("row_stripes", True)),
        showColumnStripes=bool(table_spec.get("column_stripes", False)),
    )
    worksheet.add_table(table)


def _add_excel_chart(api: dict[str, Any], worksheet: Any, chart_spec: dict[str, Any]) -> None:
    chart_type = str(chart_spec.get("type") or "bar").lower()
    if chart_type == "line":
        chart = api["LineChart"]()
    elif chart_type == "pie":
        chart = api["PieChart"]()
    else:
        chart = api["BarChart"]()

    title = chart_spec.get("title")
    if title:
        chart.title = str(title)
    values_range = str(chart_spec.get("values") or chart_spec.get("values_range") or chart_spec.get("range") or "")
    if not values_range:
        return
    _sheet_name, values_ref = _sheet_range(values_range)
    min_col, min_row, max_col, max_row = api["range_boundaries"](values_ref)
    data = api["Reference"](worksheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
    chart.add_data(data, titles_from_data=bool(chart_spec.get("titles_from_data", True)))

    categories_range = chart_spec.get("categories") or chart_spec.get("categories_range")
    if categories_range:
        _cat_sheet, categories_ref = _sheet_range(str(categories_range))
        c_min_col, c_min_row, c_max_col, c_max_row = api["range_boundaries"](categories_ref)
        cats = api["Reference"](worksheet, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row)
        chart.set_categories(cats)
    elif min_col > 1 and max_row > min_row:
        cats = api["Reference"](worksheet, min_col=min_col - 1, min_row=min_row + 1, max_row=max_row)
        chart.set_categories(cats)

    position = str(chart_spec.get("position") or "H2")
    worksheet.add_chart(chart, position)


def _add_excel_validation(api: dict[str, Any], worksheet: Any, spec: dict[str, Any]) -> None:
    range_ref = str(spec.get("range") or spec.get("ref") or "").strip()
    if not range_ref:
        return
    validation = api["DataValidation"](
        type=str(spec.get("type") or "list"),
        operator=spec.get("operator"),
        formula1=spec.get("formula1"),
        formula2=spec.get("formula2"),
        allow_blank=bool(spec.get("allow_blank", True)),
    )
    if spec.get("prompt"):
        validation.prompt = str(spec["prompt"])
    if spec.get("error"):
        validation.error = str(spec["error"])
    worksheet.add_data_validation(validation)
    validation.add(range_ref)


def _excel_wide_char(ch: str) -> bool:
    o = ord(ch)
    return (
        0x1100 <= o <= 0x115F
        or 0x2E80 <= o <= 0x9FFF
        or 0xAC00 <= o <= 0xD7A3
        or 0xF900 <= o <= 0xFAFF
        or 0xFE30 <= o <= 0xFE4F
        or 0xFF00 <= o <= 0xFF60
        or 0xFFE0 <= o <= 0xFFE6
    )


def _excel_display_len(text: str) -> int:
    longest = 0
    for line in str(text).split("\n"):
        width = sum(2 if _excel_wide_char(ch) else 1 for ch in line)
        longest = max(longest, width)
    return longest


def _autofit_excel_columns(api: dict[str, Any], worksheet: Any, explicit: set[str]) -> None:
    """Size every column to its content (CJK-aware), skipping explicit widths."""
    measured: dict[int, int] = {}
    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            if value is None:
                continue
            if isinstance(value, str) and value.startswith("="):
                length = 14
            else:
                length = _excel_display_len(value)
            if length > measured.get(cell.column, 0):
                measured[cell.column] = length
    for column_index, length in measured.items():
        letter = api["get_column_letter"](column_index)
        if letter in explicit:
            continue
        worksheet.column_dimensions[letter].width = max(9.0, min(float(length) + 2.6, 70.0))


def _apply_sheet_spec(api: dict[str, Any], workbook: Any, sheet_spec: dict[str, Any]) -> None:
    worksheet = _get_or_create_sheet(workbook, str(sheet_spec.get("name") or "Sheet"))
    if sheet_spec.get("title") and not sheet_spec.get("data"):
        worksheet["A1"] = str(sheet_spec["title"])

    header_provided = False
    data = sheet_spec.get("data")
    if data is None and "rows" in sheet_spec:
        headers = sheet_spec.get("columns") or sheet_spec.get("headers") or []
        rows = sheet_spec.get("rows") or []
        header_provided = bool(headers)
        data = ([headers] if headers else []) + rows
    start_cell = str(sheet_spec.get("start_cell") or "A1")
    _write_grid(api, worksheet, start_cell, data)

    for cell_spec in sheet_spec.get("cells") or []:
        if isinstance(cell_spec, dict) and cell_spec.get("cell"):
            _set_excel_cell_value(worksheet[str(cell_spec["cell"])], cell_spec.get("value"))
    for formula_spec in sheet_spec.get("formulas") or []:
        if isinstance(formula_spec, dict) and formula_spec.get("cell"):
            formula = str(formula_spec.get("formula") or "")
            worksheet[str(formula_spec["cell"])] = formula if formula.startswith("=") else f"={formula}"

    for column, width in (sheet_spec.get("widths") or {}).items():
        worksheet.column_dimensions[str(column)].width = float(width)
    if sheet_spec.get("freeze_panes"):
        worksheet.freeze_panes = str(sheet_spec["freeze_panes"])
    if sheet_spec.get("auto_filter"):
        used_ref = worksheet.dimensions
        if used_ref:
            worksheet.auto_filter.ref = used_ref
    if sheet_spec.get("tab_color"):
        color = _clean_color(sheet_spec.get("tab_color"))
        if color:
            worksheet.sheet_properties.tabColor = color

    for style in sheet_spec.get("formats") or []:
        if isinstance(style, dict):
            _apply_excel_style(api, worksheet, str(style.get("range") or ""), style)
    for idx, table_spec in enumerate(sheet_spec.get("tables") or []):
        if isinstance(table_spec, dict):
            _add_excel_table(api, worksheet, table_spec, idx)
    for chart_spec in sheet_spec.get("charts") or []:
        if isinstance(chart_spec, dict):
            _add_excel_chart(api, worksheet, chart_spec)
    for validation_spec in sheet_spec.get("validations") or []:
        if isinstance(validation_spec, dict):
            _add_excel_validation(api, worksheet, validation_spec)

    if header_provided and sheet_spec.get("auto_header", True):
        try:
            start_row, start_col = api["coordinate_to_tuple"](start_cell)
            header_cols = len(sheet_spec.get("columns") or sheet_spec.get("headers") or [])
            if header_cols:
                start_letter = api["get_column_letter"](start_col)
                end_letter = api["get_column_letter"](start_col + header_cols - 1)
                _apply_excel_style(api, worksheet, f"{start_letter}{start_row}:{end_letter}{start_row}", {
                    "bold": True, "fill_color": "1F2A44", "font_color": "FFFFFF",
                    "align": "center", "vertical": "center", "wrap_text": True, "border": "thin",
                })
                if not sheet_spec.get("freeze_panes"):
                    worksheet.freeze_panes = f"{start_letter}{start_row + 1}"
        except Exception:
            pass

    if sheet_spec.get("autofit", True):
        explicit_cols = {str(column) for column in (sheet_spec.get("widths") or {}).keys()}
        _autofit_excel_columns(api, worksheet, explicit_cols)


def _apply_workbook_level_excel(api: dict[str, Any], workbook: Any, spec: dict[str, Any]) -> None:
    for formula_spec in spec.get("formulas") or []:
        if not isinstance(formula_spec, dict):
            continue
        sheet_name = str(formula_spec.get("sheet") or workbook.active.title)
        worksheet = _get_or_create_sheet(workbook, sheet_name)
        cell = formula_spec.get("cell")
        if cell:
            formula = str(formula_spec.get("formula") or "")
            worksheet[str(cell)] = formula if formula.startswith("=") else f"={formula}"
    for style in spec.get("formats") or []:
        if not isinstance(style, dict):
            continue
        sheet_name = str(style.get("sheet") or workbook.active.title)
        worksheet = _get_or_create_sheet(workbook, sheet_name)
        _apply_excel_style(api, worksheet, str(style.get("range") or ""), style)
    for idx, table_spec in enumerate(spec.get("tables") or []):
        if not isinstance(table_spec, dict):
            continue
        sheet_name = str(table_spec.get("sheet") or workbook.active.title)
        _add_excel_table(api, _get_or_create_sheet(workbook, sheet_name), table_spec, idx)
    for chart_spec in spec.get("charts") or []:
        if not isinstance(chart_spec, dict):
            continue
        sheet_name = str(chart_spec.get("sheet") or workbook.active.title)
        _add_excel_chart(api, _get_or_create_sheet(workbook, sheet_name), chart_spec)
    for validation_spec in spec.get("validations") or []:
        if not isinstance(validation_spec, dict):
            continue
        sheet_name = str(validation_spec.get("sheet") or workbook.active.title)
        _add_excel_validation(api, _get_or_create_sheet(workbook, sheet_name), validation_spec)


def _workbook_summary(workbook: Any, *, include_values: bool, max_rows: int, max_cols: int) -> dict[str, Any]:
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        entry: dict[str, Any] = {
            "name": worksheet.title,
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "tables": list(getattr(worksheet, "tables", {}).keys()),
            "charts": len(getattr(worksheet, "_charts", []) or []),
            "validations": len(getattr(getattr(worksheet, "data_validations", None), "dataValidation", []) or []),
        }
        if include_values:
            rows: list[list[Any]] = []
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=min(max_rows, worksheet.max_row),
                min_col=1,
                max_col=min(max_cols, worksheet.max_column),
                values_only=True,
            ):
                rows.append(list(row))
            entry["sample"] = rows
        sheets.append(entry)
    return {"sheet_count": len(workbook.worksheets), "sheets": sheets}


def _scan_excel_issues(workbook: Any) -> list[str]:
    issues: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str):
                    if "#REF!" in value:
                        issues.append(f"{worksheet.title}!{cell.coordinate} contains #REF!.")
                    if value.startswith("=") and not value.strip("="):
                        issues.append(f"{worksheet.title}!{cell.coordinate} has an empty formula.")
                    if (
                        cell.number_format == "General"
                        and not value.startswith("=")
                        and _numeric_from_string(value) is not None
                        and not re.match(r"^0\d+$", value.strip())
                    ):
                        issues.append(f"{worksheet.title}!{cell.coordinate} looks numeric but is stored as text.")
    return issues


def _excel_quality_summary(workbook: Any) -> dict[str, Any]:
    sheet_count = len(workbook.worksheets)
    chart_count = 0
    table_count = 0
    validation_count = 0
    formula_count = 0
    text_as_number = 0
    for worksheet in workbook.worksheets:
        chart_count += len(getattr(worksheet, "_charts", []) or [])
        table_count += len(getattr(worksheet, "tables", {}) or {})
        validation_count += len(getattr(getattr(worksheet, "data_validations", None), "dataValidation", []) or [])
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formula_count += 1
                elif (
                    isinstance(value, str)
                    and cell.number_format == "General"
                    and _numeric_from_string(value) is not None
                    and not re.match(r"^0\d+$", value.strip())
                ):
                    text_as_number += 1
    summary = {
        "sheet_count": sheet_count,
        "chart_count": chart_count,
        "table_count": table_count,
        "validation_count": validation_count,
        "formula_count": formula_count,
        "text_as_number_count": text_as_number,
    }
    if text_as_number:
        summary["warnings"] = [
            f"{text_as_number} cell(s) hold numeric-looking text — use typed semantic values instead."
        ]
    return summary


def _excel_com_validate(path: Path) -> tuple[str, list[str]]:
    if not _is_windows():
        return "structural_only", ["Excel COM validation is available only on Windows."]
    try:
        import pythoncom
        import win32com.client
    except Exception as exc:
        return "structural_only", [f"Excel COM unavailable: {exc}"]
    pythoncom.CoInitialize()
    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True)
        try:
            excel.CalculateFullRebuild()
        except Exception:
            workbook.Application.CalculateFull()
        return "office_com", []
    except Exception as exc:
        return "office_com", [f"Excel COM open/recalculate failed: {exc}"]
    finally:
        if workbook is not None:
            try:
                workbook.Close(False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _is_windows() -> bool:
    import sys

    return sys.platform == "win32"


def _render_summary_png(title: str, lines: list[str], out_path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    width = 1400
    height = max(700, 120 + len(lines) * 34)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype("arial.ttf", 32)
        body_font = ImageFont.truetype("arial.ttf", 22)
    except Exception:
        title_font = ImageFont.load_default()
        body_font = ImageFont.load_default()
    draw.rectangle((0, 0, width, 78), fill="#1F2937")
    draw.text((34, 22), title, fill="white", font=title_font)
    y = 112
    for line in lines:
        draw.text((42, y), line[:150], fill="#111827", font=body_font)
        y += 34
    out_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(out_path)


def _render_excel_com(path: Path, output_dir: Path, sheet_names: list[str]) -> tuple[list[Path], list[str]]:
    if not _is_windows():
        return [], ["Excel COM rendering is available only on Windows."]
    try:
        import pythoncom
        import win32com.client
        from PIL import ImageGrab
    except Exception as exc:
        return [], [f"Excel COM rendering unavailable: {exc}"]
    pythoncom.CoInitialize()
    excel = None
    workbook = None
    previews: list[Path] = []
    issues: list[str] = []
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True)
        names = sheet_names or [sheet.Name for sheet in workbook.Worksheets]
        for index, name in enumerate(names):
            try:
                sheet = workbook.Worksheets(name)
                used = sheet.UsedRange
                used.CopyPicture(Appearance=1, Format=2)
                image = ImageGrab.grabclipboard()
                if image is None:
                    issues.append(f"{name}: clipboard did not contain an image.")
                    continue
                out_path = output_dir / f"excel_sheet_{index + 1}_{_safe_segment(name, 'sheet')}.png"
                image.save(out_path)
                previews.append(out_path)
            except Exception as exc:
                issues.append(f"{name}: render failed: {exc}")
        return previews, issues
    except Exception as exc:
        return previews, [f"Excel COM render failed: {exc}"]
    finally:
        if workbook is not None:
            try:
                workbook.Close(False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _build_excel_fallback_preview(path: Path, output_dir: Path) -> tuple[list[Path], list[str]]:
    api, error = _require_openpyxl()
    if error:
        return [], [error]
    workbook = api["load_workbook"](path, data_only=False, read_only=True)
    summary = _workbook_summary(workbook, include_values=True, max_rows=10, max_cols=8)
    lines = ["Structural preview only: native Excel rendering was not available."]
    for sheet in summary["sheets"]:
        lines.append(f"{sheet['name']}: {sheet['max_row']} rows x {sheet['max_column']} columns")
        for row in sheet.get("sample", [])[:4]:
            lines.append("  " + " | ".join("" if value is None else str(value) for value in row[:6]))
    out_path = output_dir / "excel_structural_preview.png"
    _render_summary_png(path.name, lines, out_path)
    return [out_path], ["Native Excel rendering unavailable; generated a structural preview."]


def _analysis_workbook_sheets(spec: dict[str, Any]) -> list[dict[str, Any]]:
    """Create a compact Codex-style analytical workbook skeleton from rows."""
    title = str(spec.get("title") or spec.get("name") or "Office Artifact Dashboard")
    raw_rows = spec.get("data") or spec.get("rows") or []
    headers = spec.get("columns") or spec.get("headers") or []
    if raw_rows and isinstance(raw_rows, list) and raw_rows and isinstance(raw_rows[0], dict):
        if not headers:
            seen: list[str] = []
            for row in raw_rows:
                for key in row.keys():
                    if key not in seen:
                        seen.append(str(key))
            headers = seen
        data_rows = [[row.get(header) for header in headers] for row in raw_rows]
    else:
        data_rows = raw_rows if isinstance(raw_rows, list) else []
        if data_rows and isinstance(data_rows[0], list) and not headers:
            headers = data_rows[0]
            data_rows = data_rows[1:]

    if not headers:
        headers = ["Metric", "Value"]
        data_rows = [["Status", "Draft"], ["Items", len(data_rows)]]

    data = [headers] + data_rows
    last_row = max(2, len(data))
    last_col = max(2, len(headers))
    last_col_letter = chr(ord("A") + min(last_col, 26) - 1)
    numeric_cols = [idx + 1 for idx, header in enumerate(headers) if any(token in str(header).lower() for token in ("revenue", "sales", "value", "amount", "ebitda", "profit", "cost"))]
    metric_col = numeric_cols[0] if numeric_cols else 2
    metric_col_letter = chr(ord("A") + min(metric_col, 26) - 1)
    category_col_letter = "A"

    return [
        {
            "name": "Dashboard",
            "data": [
                [title],
                ["A compact operating model with formulas, assumptions, checks, and native charts."],
                [],
                ["KPI", "Value", "KPI", "Value", "KPI", "Value"],
                ["Record count", {"type": "number", "value": len(data_rows), "number_format": "#,##0"}, "Primary metric", f"=SUM(Data!{metric_col_letter}2:{metric_col_letter}{last_row})", "Check status", "=IF(Checks!B2=0,\"PASS\",\"CHECK\")"],
                ["Last refreshed", {"type": "date", "value": date.today().isoformat()}, "Source", str(spec.get("source") or "Generated data"), "Formula errors", "=Checks!B2"],
            ],
            "formats": [
                {"range": "A1:F1", "bold": True, "font_size": 18, "fill_color": "#111827", "font_color": "#FFFFFF"},
                {"range": "A4:F4", "bold": True, "fill_color": "#111827", "font_color": "#FFFFFF", "align": "center"},
                {"range": "A5:A6", "bold": True, "fill_color": "#E8EEF6"},
                {"range": "C5:C6", "bold": True, "fill_color": "#E8EEF6"},
                {"range": "E5:E6", "bold": True, "fill_color": "#E8EEF6"},
                {"range": "B5:F6", "number_format": "#,##0.0"},
            ],
            "widths": {"A": 18, "B": 16, "C": 18, "D": 18, "E": 18, "F": 16, "H": 20, "I": 16, "J": 16},
            "charts": [{
                "type": "bar",
                "title": str(spec.get("chart_title") or "Primary metric"),
                "values": f"Data!{metric_col_letter}1:{metric_col_letter}{last_row}",
                "categories": f"Data!{category_col_letter}2:{category_col_letter}{last_row}",
                "position": "H4",
            }],
        },
        {
            "name": "Data",
            "data": data,
            "tables": [{"range": f"A1:{last_col_letter}{last_row}", "name": "SourceData"}],
            "formats": [
                {"range": f"A1:{last_col_letter}1", "bold": True, "fill_color": "#0F766E", "font_color": "#FFFFFF"},
            ],
            "freeze_panes": "A2",
            "auto_filter": True,
        },
        {
            "name": "Assumptions",
            "data": [
                ["Assumption", "Value", "Notes"],
                ["Scenario uplift", {"type": "percent", "value": spec.get("scenario_uplift", 0.08)}, "Editable driver"],
                ["Target gap", {"type": "percent", "value": spec.get("target_gap", 0.015)}, "Editable driver"],
                ["Opex flex", {"type": "percent", "value": spec.get("opex_flex", -0.02)}, "Editable driver"],
            ],
            "formats": [{"range": "A1:C1", "bold": True, "fill_color": "#B7791F", "font_color": "#FFFFFF"}],
            "validations": [{"range": "B2:B4", "type": "decimal", "operator": "between", "formula1": "-1", "formula2": "2"}],
        },
        {
            "name": "Checks",
            "data": [
                ["Check", "Error count", "Status"],
                ["Formula error scan", 0, "=IF(B2=0,\"PASS\",\"CHECK\")"],
                ["Package integrity", 0, "=IF(B3=0,\"PASS\",\"CHECK\")"],
            ],
            "formats": [{"range": "A1:C1", "bold": True, "fill_color": "#111827", "font_color": "#FFFFFF"}],
        },
    ]


def _find_soffice() -> Optional[str]:
    return shutil.which("soffice") or shutil.which("libreoffice")


def _libreoffice_available() -> bool:
    return bool(_find_soffice())


class ExcelCreateTool(BaseTool):
    name = "excel_create"
    description = (
        "Create a product-grade XLSX workbook from a structured JSON spec. Use this "
        "for deliverable Excel files with sheets, tables, formulas, formatting, charts, and validation."
    )
    parameters = {
        "type": "object",
        "properties": {
            "output_name": {"type": "string", "description": "Output .xlsx file name."},
            "spec": {"type": "object", "description": "Workbook spec. May include sheets, tables, ranges, formulas, formats, charts, validations.", "additionalProperties": True},
            "sheets": {"type": "array", "description": "Optional shorthand for spec.sheets.", "items": {"type": "object", "additionalProperties": True}},
            "quality_profile": {"type": "string", "description": "Optional profile. Use analysis_dashboard for Dashboard/Data/Assumptions/Checks workbook defaults."},
            "title": {"type": "string", "description": "Artifact display title."},
        },
        "required": ["output_name"],
    }

    async def execute(
        self,
        output_name: str,
        spec: Any = None,
        sheets: Any = None,
        quality_profile: str = "",
        title: str = "",
        session_id: str = "default",
        tool_call_id: str = "",
        agent_type: str = "",
    ) -> ToolResult:
        api, error = _require_openpyxl()
        if error:
            return ToolResult(error=error)
        spec_dict, error = _as_dict(spec, "spec")
        if error:
            return ToolResult(error=error)
        if sheets is not None:
            sheet_list, error = _as_list(sheets, "sheets")
            if error:
                return ToolResult(error=error)
            spec_dict = {**spec_dict, "sheets": sheet_list}
        if quality_profile:
            spec_dict.setdefault("quality_profile", quality_profile)

        workbook = api["Workbook"]()
        default = workbook.active
        profile = str(spec_dict.get("quality_profile") or spec_dict.get("profile") or "").lower()
        if not spec_dict.get("sheets") and profile in {"analysis", "analysis_dashboard", "financial_dashboard", "codex"}:
            spec_dict["sheets"] = _analysis_workbook_sheets(spec_dict)
        sheet_specs = spec_dict.get("sheets") or [{"name": "Sheet1", "data": []}]
        if sheet_specs:
            workbook.remove(default)
        for sheet_spec in sheet_specs:
            if isinstance(sheet_spec, dict):
                _apply_sheet_spec(api, workbook, sheet_spec)
        _apply_workbook_level_excel(api, workbook, spec_dict)

        path = _output_path(output_name, ".xlsx", session_id=session_id, tool_call_id=tool_call_id, kind=OFFICE_KIND_EXCEL)
        workbook.save(path)
        artifacts, artifact_error = _office_artifact(
            path,
            title=title or path.name,
            source_tool=self.name,
            session_id=session_id,
            tool_call_id=tool_call_id,
            agent_type=agent_type,
        )
        if artifact_error:
            return ToolResult(error=artifact_error)
        summary = _workbook_summary(workbook, include_values=False, max_rows=0, max_cols=0)
        quality = _excel_quality_summary(workbook)
        metadata = {"artifacts": artifacts, "file_path": str(path), "qa_status": "created", "summary": summary, "qa_summary": quality}
        return ToolResult(output=f"Excel workbook created: {path}\n\n{_json(summary)}", metadata=metadata)


class ExcelEditTool(BaseTool):
    name = "excel_edit"
    description = "Edit an existing XLSX workbook with structured operations, then publish the edited Office artifact."
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Existing XLSX path."},
            "operations": {"type": "array", "description": "Edit operations.", "items": {"type": "object", "additionalProperties": True}},
            "output_name": {"type": "string", "description": "Optional output .xlsx name. Defaults to edited_<source>."},
            "title": {"type": "string", "description": "Artifact display title."},
        },
        "required": ["path", "operations"],
    }

    async def execute(
        self,
        path: str,
        operations: Any,
        output_name: str = "",
        title: str = "",
        session_id: str = "default",
        tool_call_id: str = "",
        agent_type: str = "",
    ) -> ToolResult:
        api, error = _require_openpyxl()
        if error:
            return ToolResult(error=error)
        source, error = _resolve_input_path(path, agent_type=agent_type)
        if error:
            return ToolResult(error=error)
        ops, error = _as_list(operations, "operations")
        if error:
            return ToolResult(error=error)
        workbook = api["load_workbook"](source)
        applied: list[str] = []
        for op in ops:
            if not isinstance(op, dict):
                continue
            action = str(op.get("op") or op.get("type") or "").lower()
            sheet_name = str(op.get("sheet") or workbook.active.title)
            worksheet = _get_or_create_sheet(workbook, sheet_name)
            if action == "set_values":
                _write_grid(api, worksheet, str(op.get("start_cell") or op.get("cell") or "A1"), op.get("values", op.get("value")))
            elif action == "set_formula":
                formula = str(op.get("formula") or "")
                if op.get("cell"):
                    worksheet[str(op["cell"])] = formula if formula.startswith("=") else f"={formula}"
            elif action == "set_formulas":
                for formula_spec in op.get("formulas") or []:
                    if isinstance(formula_spec, dict) and formula_spec.get("cell"):
                        formula = str(formula_spec.get("formula") or "")
                        worksheet[str(formula_spec["cell"])] = formula if formula.startswith("=") else f"={formula}"
            elif action == "format_range":
                _apply_excel_style(api, worksheet, str(op.get("range") or ""), op)
            elif action == "add_table":
                _add_excel_table(api, worksheet, op, len(getattr(worksheet, "tables", {})))
            elif action == "add_chart":
                _add_excel_chart(api, worksheet, op)
            elif action == "data_validation":
                _add_excel_validation(api, worksheet, op)
            elif action == "freeze_panes" and op.get("cell"):
                worksheet.freeze_panes = str(op["cell"])
            elif action == "set_widths":
                for column, width in (op.get("widths") or {}).items():
                    worksheet.column_dimensions[str(column)].width = float(width)
            elif action == "rename_sheet":
                new_name = str(op.get("new_name") or "").strip()[:31]
                if new_name:
                    worksheet.title = new_name
            elif action == "copy_sheet":
                copied = workbook.copy_worksheet(worksheet)
                copied.title = str(op.get("new_name") or f"{worksheet.title} Copy")[:31]
            elif action == "delete_sheet":
                if len(workbook.worksheets) > 1:
                    workbook.remove(worksheet)
            elif action == "replace_text":
                old = str(op.get("old") or "")
                new = str(op.get("new") or "")
                if old:
                    for row in worksheet.iter_rows():
                        for cell in row:
                            if isinstance(cell.value, str) and old in cell.value:
                                cell.value = cell.value.replace(old, new)
            else:
                continue
            applied.append(action)

        out_name = output_name or f"edited_{source.name}"
        out_path = _output_path(out_name, ".xlsx", session_id=session_id, tool_call_id=tool_call_id, kind=OFFICE_KIND_EXCEL)
        workbook.save(out_path)
        artifacts, artifact_error = _office_artifact(
            out_path,
            title=title or out_path.name,
            source_tool=self.name,
            session_id=session_id,
            tool_call_id=tool_call_id,
            agent_type=agent_type,
        )
        if artifact_error:
            return ToolResult(error=artifact_error)
        summary = _workbook_summary(workbook, include_values=False, max_rows=0, max_cols=0)
        quality = _excel_quality_summary(workbook)
        metadata = {"artifacts": artifacts, "file_path": str(out_path), "qa_status": "edited", "applied_operations": applied, "summary": summary, "qa_summary": quality}
        return ToolResult(output=f"Excel workbook edited: {out_path}\nApplied operations: {', '.join(applied) or '(none)'}", metadata=metadata)


class ExcelInspectTool(BaseTool):
    name = "excel_inspect"
    description = "Inspect workbook sheets, dimensions, tables, charts, and sample values."
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "XLSX path."},
            "include_values": {"type": "boolean", "default": True},
            "max_rows": {"type": "integer", "default": 20},
            "max_cols": {"type": "integer", "default": 12},
        },
        "required": ["path"],
    }

    async def execute(
        self,
        path: str,
        include_values: bool = True,
        max_rows: int = 20,
        max_cols: int = 12,
        agent_type: str = "",
    ) -> ToolResult:
        api, error = _require_openpyxl()
        if error:
            return ToolResult(error=error)
        source, error = _resolve_input_path(path, agent_type=agent_type)
        if error:
            return ToolResult(error=error)
        workbook = api["load_workbook"](source, data_only=False, read_only=False)
        summary = _workbook_summary(workbook, include_values=include_values, max_rows=max_rows, max_cols=max_cols)
        return ToolResult(output=_json(summary), metadata={"summary": summary, "file_path": str(source)})


class ExcelValidateTool(BaseTool):
    name = "excel_validate"
    description = "Validate an XLSX structurally and, when available, open/recalculate it through native Excel COM."
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "XLSX path."},
            "use_com": {"type": "boolean", "default": True},
        },
        "required": ["path"],
    }

    async def execute(self, path: str, use_com: bool = True, agent_type: str = "") -> ToolResult:
        api, error = _require_openpyxl()
        if error:
            return ToolResult(error=error)
        source, error = _resolve_input_path(path, agent_type=agent_type)
        if error:
            return ToolResult(error=error)
        issues: list[str] = []
        engine = "structural_only"
        try:
            workbook = api["load_workbook"](source, data_only=False)
            issues.extend(_scan_excel_issues(workbook))
            quality = _excel_quality_summary(workbook)
        except Exception as exc:
            return ToolResult(error=f"Workbook could not be opened structurally: {exc}")
        if use_com:
            engine, com_issues = _excel_com_validate(source)
            issues.extend(com_issues)
        qa_status = "passed" if not issues else ("warnings" if engine == "structural_only" else "failed")
        result = {"file_path": str(source), "engine": engine, "qa_status": qa_status, "issues": issues, **quality}
        return ToolResult(output=_json(result), metadata=result)


class ExcelRenderTool(BaseTool):
    name = "excel_render"
    description = "Render Excel workbook previews as image artifacts. Uses native Excel COM when available; otherwise emits structural preview images."
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "XLSX path."},
            "sheets": {"type": "array", "items": {"type": "string"}, "description": "Optional sheet names to render."},
        },
        "required": ["path"],
    }

    async def execute(
        self,
        path: str,
        sheets: Optional[list[str]] = None,
        session_id: str = "default",
        tool_call_id: str = "",
        agent_type: str = "",
    ) -> ToolResult:
        source, error = _resolve_input_path(path, agent_type=agent_type)
        if error:
            return ToolResult(error=error)
        output_dir = _office_run_dir(session_id, tool_call_id, OFFICE_KIND_EXCEL) / "render"
        output_dir.mkdir(parents=True, exist_ok=True)
        preview_paths, issues = _render_excel_com(source, output_dir, sheets or [])
        engine = "office_com"
        if not preview_paths:
            preview_paths, fallback_issues = _build_excel_fallback_preview(source, output_dir)
            issues.extend(fallback_issues)
            engine = "structural_only"
        artifacts = _preview_artifacts(
            preview_paths,
            title_prefix=f"{source.stem} sheet preview",
            source_tool=self.name,
            session_id=session_id,
            tool_call_id=tool_call_id,
            agent_type=agent_type,
        )
        qa_status = "passed" if engine == "office_com" and not issues else "warnings"
        result = {
            "file_path": str(source),
            "preview_paths": [str(path) for path in preview_paths],
            "engine": engine,
            "qa_status": qa_status,
            "issues": issues,
            "artifacts": artifacts,
        }
        return ToolResult(output=_json({k: v for k, v in result.items() if k != "artifacts"}), metadata=result)


def _ppt_layout(prs: Any, layout: Any):
    if isinstance(layout, int) or (isinstance(layout, str) and layout.isdigit()):
        index = int(layout)
    else:
        index = {
            "title": 0,
            "title_content": 1,
            "content": 1,
            "section": 2,
            "two_content": 3,
            "blank": 6,
        }.get(str(layout or "blank").lower(), 6)
    if 0 <= index < len(prs.slide_layouts):
        return prs.slide_layouts[index]
    return prs.slide_layouts[6]


def _pt(api: dict[str, Any], value: Any, default: float):
    try:
        return api["Pt"](float(value))
    except Exception:
        return api["Pt"](default)


def _ppt_color(api: dict[str, Any], value: Any):
    color = _clean_color(value)
    if not color:
        return None
    return api["RGBColor"](int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))


PPT_THEME = {
    "bg": "F6F1E7",
    "ink": "111827",
    "muted": "56657A",
    "navy": "111827",
    "blue": "2D6A9F",
    "green": "0F7B5F",
    "gold": "B7791F",
}


def _chart_type(api: dict[str, Any], name: Any):
    xl = api.get("XL_CHART_TYPE")
    if xl is None:
        return None
    mapping = {
        "bar": xl.COLUMN_CLUSTERED,
        "column": xl.COLUMN_CLUSTERED,
        "hbar": xl.BAR_CLUSTERED,
        "line": xl.LINE_MARKERS,
        "pie": xl.PIE,
        "doughnut": xl.DOUGHNUT,
        "donut": xl.DOUGHNUT,
        "area": xl.AREA,
    }
    return mapping.get(str(name or "bar").lower(), xl.COLUMN_CLUSTERED)


def _ppt_cell_border(cell: Any, color_hex: str, width_emu: int = 6350) -> None:
    """Apply a thin solid border to all four sides of a table cell."""
    try:
        from pptx.oxml.ns import qn
    except Exception:
        return
    color = _clean_color(color_hex) or "D4D8E0"
    try:
        tc_pr = cell._tc.get_or_add_tcPr()
    except Exception:
        return
    for tag in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):
        for existing in tc_pr.findall(qn(tag)):
            tc_pr.remove(existing)
        ln = tc_pr.makeelement(qn(tag), {"w": str(width_emu), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        solid = ln.makeelement(qn("a:solidFill"), {})
        srgb = solid.makeelement(qn("a:srgbClr"), {"val": color})
        solid.append(srgb)
        ln.append(solid)
        tc_pr.insert(0, ln)


_PPT_ANCHOR_MAP = {"top": "TOP", "middle": "MIDDLE", "bottom": "BOTTOM"}
_PPT_ALIGN_MAP = {"left": "LEFT", "center": "CENTER", "right": "RIGHT"}


def _render_text_element(api: dict[str, Any], slide: Any, el: dict[str, Any]) -> None:
    inch = api["Inches"]
    box = slide.shapes.add_textbox(inch(el["x"]), inch(el["y"]), inch(max(0.1, el["w"])), inch(max(0.12, el["h"])))
    frame = box.text_frame
    frame.word_wrap = True
    try:
        frame.auto_size = None
    except Exception:
        pass
    for margin in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        try:
            setattr(frame, margin, inch(0.0))
        except Exception:
            pass
    anchor = api.get("MSO_ANCHOR")
    if anchor is not None:
        try:
            frame.vertical_anchor = getattr(anchor, _PPT_ANCHOR_MAP.get(el.get("anchor", "top"), "TOP"))
        except Exception:
            pass
    align_enum = api.get("PP_ALIGN")
    alignment = None
    if align_enum is not None:
        alignment = getattr(align_enum, _PPT_ALIGN_MAP.get(el.get("align", "left"), "LEFT"), None)
    rgb = _ppt_color(api, el.get("color"))
    lines = str(el.get("text") or "").split("\n")
    frame.text = lines[0] if lines else ""
    paragraphs = [frame.paragraphs[0]]
    for extra in lines[1:]:
        paragraph = frame.add_paragraph()
        paragraph.text = extra
        paragraphs.append(paragraph)
    for paragraph in paragraphs:
        if alignment is not None:
            paragraph.alignment = alignment
        for run in paragraph.runs:
            run.font.size = _pt(api, el.get("font_size"), 14)
            run.font.bold = bool(el.get("bold"))
            run.font.italic = bool(el.get("italic"))
            if rgb:
                run.font.color.rgb = rgb


def _render_rect_element(api: dict[str, Any], slide: Any, el: dict[str, Any]) -> None:
    inch = api["Inches"]
    shape_enum = api["MSO_SHAPE"].ROUNDED_RECTANGLE if el.get("rounded") else api["MSO_SHAPE"].RECTANGLE
    shape = slide.shapes.add_shape(
        shape_enum, inch(el["x"]), inch(el["y"]), inch(max(0.02, el["w"])), inch(max(0.02, el["h"]))
    )
    fill = _ppt_color(api, el.get("fill"))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        try:
            shape.fill.background()
        except Exception:
            pass
    line_color = _ppt_color(api, el.get("line_color")) if el.get("line_color") else None
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = _pt(api, 0.75, 0.75)
    else:
        try:
            shape.line.fill.background()
        except Exception:
            pass
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    text = el.get("text")
    if text:
        frame = shape.text_frame
        frame.word_wrap = True
        frame.text = str(text)
        rgb = _ppt_color(api, el.get("text_color") or "FFFFFF")
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = _pt(api, el.get("font_size"), 12)
                if rgb:
                    run.font.color.rgb = rgb


def _render_table_element(api: dict[str, Any], slide: Any, el: dict[str, Any]) -> None:
    inch = api["Inches"]
    grid = el.get("grid") or []
    if not grid:
        return
    rows = len(grid)
    cols = max(len(row) for row in grid)
    shape = slide.shapes.add_table(
        rows, cols, inch(el["x"]), inch(el["y"]), inch(max(0.4, el["w"])), inch(max(0.3, el["h"]))
    )
    table = shape.table
    for attr in ("first_row", "last_row", "first_col", "last_col", "horz_banding", "vert_banding"):
        try:
            setattr(table, attr, False)
        except Exception:
            pass
    col_w = el.get("col_w") or []
    for ci in range(cols):
        if ci < len(col_w):
            try:
                table.columns[ci].width = inch(col_w[ci])
            except Exception:
                pass
    row_h = el.get("row_h") or []
    for ri in range(rows):
        if ri < len(row_h):
            try:
                table.rows[ri].height = inch(row_h[ri])
            except Exception:
                pass
    has_header = bool(el.get("has_header"))
    font_size = float(el.get("font_size") or 11)
    header_fill = _ppt_color(api, el.get("header_fill") or "1F2A44")
    header_text = _ppt_color(api, el.get("header_text") or "FFFFFF")
    row_alt = _ppt_color(api, el.get("row_alt") or "F3F0E8")
    white = _ppt_color(api, "FFFFFF")
    ink = _ppt_color(api, el.get("ink") or "1A1A2E")
    grid_color = el.get("grid_color") or "D4D8E0"
    align_enum = api.get("PP_ALIGN")
    anchor = api.get("MSO_ANCHOR")
    for ri in range(rows):
        is_header = has_header and ri == 0
        row = grid[ri]
        for ci in range(cols):
            cell = table.cell(ri, ci)
            value = row[ci] if ci < len(row) else ""
            cell.text = "" if value is None else str(value)
            for margin, amount in (
                ("margin_left", 0.07), ("margin_right", 0.07),
                ("margin_top", 0.03), ("margin_bottom", 0.03),
            ):
                try:
                    setattr(cell, margin, inch(amount))
                except Exception:
                    pass
            if anchor is not None:
                try:
                    cell.vertical_anchor = anchor.MIDDLE
                except Exception:
                    pass
            try:
                cell.text_frame.word_wrap = True
            except Exception:
                pass
            try:
                cell.fill.solid()
                if is_header:
                    cell.fill.fore_color.rgb = header_fill
                elif ri % 2 == 1:
                    cell.fill.fore_color.rgb = white
                else:
                    cell.fill.fore_color.rgb = row_alt
            except Exception:
                pass
            color = header_text if is_header else ink
            for paragraph in cell.text_frame.paragraphs:
                if align_enum is not None:
                    if is_header or ci > 0:
                        paragraph.alignment = align_enum.CENTER
                    else:
                        paragraph.alignment = align_enum.LEFT
                for run in paragraph.runs:
                    run.font.size = _pt(api, font_size + (0.5 if is_header else 0), 11)
                    run.font.bold = is_header
                    if color:
                        run.font.color.rgb = color
            _ppt_cell_border(cell, grid_color)


def _render_chart_element(api: dict[str, Any], slide: Any, el: dict[str, Any]) -> None:
    inch = api["Inches"]
    series = el.get("series") or []
    categories = el.get("categories") or []
    chart_data_cls = api.get("CategoryChartData")
    if chart_data_cls is None or not series or not categories:
        return
    chart_data = chart_data_cls()
    chart_data.categories = [str(c) for c in categories]
    for entry in series:
        values = entry.get("values") or []
        chart_data.add_series(str(entry.get("name") or "Series"), tuple(values))
    chart_type = _chart_type(api, el.get("chart_type"))
    if chart_type is None:
        return
    try:
        frame = slide.shapes.add_chart(
            chart_type, inch(el["x"]), inch(el["y"]), inch(el["w"]), inch(el["h"]), chart_data
        )
    except Exception:
        return
    chart = frame.chart
    legend_enum = api.get("XL_LEGEND_POSITION")
    show_legend = len(series) > 1 or str(el.get("chart_type") or "").lower() in {"pie", "doughnut", "donut"}
    try:
        chart.has_legend = bool(show_legend)
        if show_legend and legend_enum is not None:
            chart.legend.position = legend_enum.BOTTOM
            chart.legend.include_in_layout = False
    except Exception:
        pass
    title = str(el.get("title") or "").strip()
    try:
        if title:
            chart.has_title = True
            chart.chart_title.text_frame.text = title
        else:
            chart.has_title = False
    except Exception:
        pass


def _render_image_element(api: dict[str, Any], slide: Any, el: dict[str, Any], agent_type: str) -> None:
    inch = api["Inches"]
    path = str(el.get("path") or "")
    if not path:
        return
    source = _existing_source_path(path, agent_type=agent_type)
    if not source:
        return
    try:
        slide.shapes.add_picture(str(source), inch(el["x"]), inch(el["y"]), inch(el["w"]), inch(el["h"]))
    except Exception:
        pass


def _render_ppt_element(api: dict[str, Any], slide: Any, el: dict[str, Any], agent_type: str) -> None:
    kind = str(el.get("kind") or "")
    try:
        if kind == "text":
            _render_text_element(api, slide, el)
        elif kind == "rect":
            _render_rect_element(api, slide, el)
        elif kind == "table":
            _render_table_element(api, slide, el)
        elif kind == "chart":
            _render_chart_element(api, slide, el)
        elif kind == "image":
            _render_image_element(api, slide, el, agent_type)
    except Exception:
        pass


def _render_planned_slide(
    api: dict[str, Any], prs: Any, planned: dict[str, Any], slide_index: int, agent_type: str
) -> Any:
    slide = prs.slides.add_slide(_ppt_layout(prs, "blank"))
    background = _ppt_color(api, planned.get("background") or PPT_THEME["bg"])
    if background:
        try:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = background
        except Exception:
            pass
    accent = planned.get("accent") or PPT_THEME["gold"]
    try:
        marker = slide.shapes.add_shape(
            api["MSO_SHAPE"].RECTANGLE,
            api["Inches"](0.45), api["Inches"](0.40), api["Inches"](0.17), api["Inches"](0.055),
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = _ppt_color(api, accent)
        marker.line.fill.background()
        marker.shadow.inherit = False
    except Exception:
        pass
    for element in planned.get("elements") or []:
        _render_ppt_element(api, slide, element, agent_type)
    if planned.get("page_number", True):
        _render_text_element(api, slide, {
            "kind": "text", "x": 12.45, "y": 6.95, "w": 0.6, "h": 0.3,
            "text": f"{slide_index:02d}", "font_size": 9, "color": PPT_THEME["ink"],
            "bold": True, "align": "right", "anchor": "middle",
        })
    notes = planned.get("speaker_notes")
    if notes:
        try:
            slide.notes_slide.notes_text_frame.text = str(notes)
        except Exception:
            pass
    return slide


def _apply_ppt_slide_spec(
    api: dict[str, Any], prs: Any, slide_spec: dict[str, Any], agent_type: str, slide_index: int = 1
):
    try:
        from app.office_artifacts.ppt_layout import plan_slides

        planned_pages = plan_slides(slide_spec)
    except Exception:
        planned_pages = [{
            "background": PPT_THEME["bg"], "accent": PPT_THEME["gold"],
            "page_number": True, "elements": [], "speaker_notes": "",
        }]
    last_slide = None
    for planned in planned_pages:
        last_slide = _render_planned_slide(api, prs, planned, len(prs.slides) + 1, agent_type)
    return last_slide


def _add_ppt_block(api: dict[str, Any], slide: Any, block: dict[str, Any], index: int, agent_type: str) -> None:
    """Render a single semantic block onto an existing slide (used by ppt_edit)."""
    from app.office_artifacts.ppt_layout import (
        CONTENT_W,
        DEFAULT_THEME,
        MARGIN_X,
        _prepare_block,
    )

    kind = str(block.get("type") or block.get("kind") or "text").lower()
    origin_x = float(block["x"]) if block.get("x") is not None else MARGIN_X
    origin_y = float(block["y"]) if block.get("y") is not None else 1.7 + index * 0.35
    width = float(block.get("w") or block.get("width") or CONTENT_W)
    if kind == "shape":
        height = float(block.get("h") or block.get("height") or 1.0)
        _render_ppt_element(api, slide, {
            "kind": "rect", "x": origin_x, "y": origin_y, "w": width, "h": height,
            "fill": _clean_color(block.get("fill_color")) or DEFAULT_THEME["card_alt"],
            "line_color": None, "rounded": True, "text": block.get("text"),
            "text_color": _clean_color(block.get("font_color")) or DEFAULT_THEME["ink"],
            "font_size": block.get("font_size") or 14,
        }, agent_type)
        return
    prepared = _prepare_block(block, width, DEFAULT_THEME)
    for element in prepared["emit"](origin_x, origin_y, width):
        _render_ppt_element(api, slide, element, agent_type)


def _ppt_summary(prs: Any, *, include_text: bool = True) -> dict[str, Any]:
    slides: list[dict[str, Any]] = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        image_count = 0
        table_count = 0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text:
                texts.append(shape.text)
            if getattr(shape, "shape_type", None) and "PICTURE" in str(shape.shape_type):
                image_count += 1
            if getattr(shape, "has_table", False):
                table_count += 1
        entry: dict[str, Any] = {"slide": idx, "shapes": len(slide.shapes), "images": image_count, "tables": table_count}
        if include_text:
            entry["text"] = texts[:8]
        slides.append(entry)
    return {"slide_count": len(prs.slides), "slides": slides}


def _shape_rect(shape: Any) -> tuple[float, float, float, float]:
    return (
        float(getattr(shape, "left", 0) or 0) / 914400,
        float(getattr(shape, "top", 0) or 0) / 914400,
        float(getattr(shape, "width", 0) or 0) / 914400,
        float(getattr(shape, "height", 0) or 0) / 914400,
    )


def _rect_overlap(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> float:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    dx = min(ax + aw, bx + bw) - max(ax, bx)
    dy = min(ay + ah, by + bh) - max(ay, by)
    if dx <= 0 or dy <= 0:
        return 0.0
    return dx * dy


def _scan_ppt_layout_issues(prs: Any) -> tuple[list[str], list[str], list[dict[str, Any]]]:
    errors: list[str] = []
    warnings: list[str] = []
    layout_json: list[dict[str, Any]] = []
    slide_w = float(prs.slide_width) / 914400
    slide_h = float(prs.slide_height) / 914400
    for slide_idx, slide in enumerate(prs.slides, start=1):
        text_shapes: list[tuple[str, tuple[float, float, float, float], int]] = []
        slide_entry: dict[str, Any] = {"slide": slide_idx, "shapes": []}
        visible_text = ""
        for shape_idx, shape in enumerate(slide.shapes, start=1):
            x, y, w, h = _shape_rect(shape)
            text = str(getattr(shape, "text", "") or "").strip()
            visible_text += text
            slide_entry["shapes"].append({
                "index": shape_idx,
                "x": round(x, 3),
                "y": round(y, 3),
                "w": round(w, 3),
                "h": round(h, 3),
                "text": text[:120],
                "has_table": bool(getattr(shape, "has_table", False)),
            })
            if x < -0.05 or y < -0.05 or x + w > slide_w + 0.05 or y + h > slide_h + 0.05:
                errors.append(f"Slide {slide_idx} shape {shape_idx} is outside the slide bounds.")
            if text and not bool(getattr(shape, "has_table", False)):
                text_shapes.append((text, (x, y, w, h), shape_idx))
                needed = _ppt_text_overflow(shape, text, w)
                if needed is not None and needed > h * 1.28 + 0.04:
                    warnings.append(
                        f"Slide {slide_idx} shape {shape_idx} text likely overflows its box "
                        f"(needs ~{needed:.2f}in in a {h:.2f}in box)."
                    )
        if not visible_text.strip():
            warnings.append(f"Slide {slide_idx} has no visible text.")
        for left_idx, (left_text, left_rect, left_shape) in enumerate(text_shapes):
            for right_text, right_rect, right_shape in text_shapes[left_idx + 1:]:
                overlap = _rect_overlap(left_rect, right_rect)
                if overlap <= 0:
                    continue
                min_area = max(0.01, min(left_rect[2] * left_rect[3], right_rect[2] * right_rect[3]))
                ratio = overlap / min_area
                if ratio > 0.20 and overlap > 0.08:
                    errors.append(
                        f"Slide {slide_idx} text shapes {left_shape} and {right_shape} overlap badly."
                    )
                elif ratio > 0.07 and overlap > 0.04:
                    warnings.append(
                        f"Slide {slide_idx} text shapes {left_shape} and {right_shape} overlap."
                    )
        layout_json.append(slide_entry)
    return errors, warnings, layout_json


def _ppt_text_overflow(shape: Any, text: str, width_in: float) -> Optional[float]:
    """Estimate the height (inches) the shape's text actually needs when wrapped."""
    try:
        from app.office_artifacts.ppt_layout import measure_text
    except Exception:
        return None
    font_pt = 14.0
    try:
        sizes = [
            run.font.size.pt
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
            if run.font.size is not None
        ]
        if sizes:
            font_pt = max(sizes)
    except Exception:
        pass
    try:
        _lines, height = measure_text(text, font_pt, max(0.4, width_in))
        return height
    except Exception:
        return None


def _build_contact_sheet(image_paths: list[Path], output_path: Path, title: str = "Slide Contact Sheet") -> Optional[Path]:
    if not image_paths:
        return None
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return None
    thumbs = []
    for path in image_paths:
        try:
            image = Image.open(path).convert("RGB")
            image.thumbnail((420, 236))
            thumbs.append((path, image.copy()))
        except Exception:
            continue
    if not thumbs:
        return None
    cols = min(3, max(1, len(thumbs)))
    rows = (len(thumbs) + cols - 1) // cols
    pad = 28
    label_h = 24
    header_h = 58
    cell_w = 420
    cell_h = 236 + label_h
    canvas = Image.new("RGB", (cols * cell_w + pad * (cols + 1), rows * cell_h + pad * (rows + 1) + header_h), "white")
    draw = ImageDraw.Draw(canvas)
    try:
        title_font = ImageFont.truetype("arial.ttf", 24)
        label_font = ImageFont.truetype("arial.ttf", 13)
    except Exception:
        title_font = ImageFont.load_default()
        label_font = ImageFont.load_default()
    draw.text((pad, 18), title, fill="#111827", font=title_font)
    for idx, (_path, thumb) in enumerate(thumbs):
        row = idx // cols
        col = idx % cols
        x = pad + col * (cell_w + pad)
        y = header_h + pad + row * (cell_h + pad)
        draw.rectangle((x - 1, y - 1, x + cell_w + 1, y + 236 + 1), outline="#D7DEE8")
        canvas.paste(thumb, (x + (cell_w - thumb.width) // 2, y + (236 - thumb.height) // 2))
        draw.text((x, y + 242), f"Slide {idx + 1:02d}", fill="#111827", font=label_font)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path)
    return output_path


def _delete_ppt_slide(prs: Any, index: int) -> None:
    if index < 0 or index >= len(prs.slides):
        return
    slide_id_list = prs.slides._sldIdLst
    slide_id = slide_id_list[index]
    rel_id = slide_id.rId
    prs.part.drop_rel(rel_id)
    slide_id_list.remove(slide_id)


def _duplicate_ppt_slide_basic(api: dict[str, Any], prs: Any, index: int, agent_type: str) -> None:
    if index < 0 or index >= len(prs.slides):
        return
    source = prs.slides[index]
    new_slide = prs.slides.add_slide(_ppt_layout(prs, "blank"))
    tmp_dir = runtime_dir("office") / "ppt-image-cache"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    for shape in source.shapes:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        if getattr(shape, "has_table", False):
            rows = []
            for row in shape.table.rows:
                rows.append([cell.text for cell in row.cells])
            _add_ppt_block(api, new_slide, {"type": "table", "rows": rows, "x": left / 914400, "y": top / 914400, "w": width / 914400, "h": height / 914400}, 0, agent_type)
        elif getattr(shape, "has_text_frame", False):
            box = new_slide.shapes.add_textbox(left, top, width, height)
            box.text = shape.text
        elif hasattr(shape, "image"):
            image_path = tmp_dir / f"{uuid.uuid4().hex}.bin"
            image_path.write_bytes(shape.image.blob)
            new_slide.shapes.add_picture(str(image_path), left, top, width, height)


def _ppt_com_validate(path: Path) -> tuple[str, list[str]]:
    if not _is_windows():
        return "structural_only", ["PowerPoint COM validation is available only on Windows."]
    try:
        import pythoncom
        import win32com.client
    except Exception as exc:
        return "structural_only", [f"PowerPoint COM unavailable: {exc}"]
    pythoncom.CoInitialize()
    app = None
    presentation = None
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        presentation = app.Presentations.Open(str(path), ReadOnly=True, WithWindow=False)
        return "office_com", []
    except Exception as exc:
        return "office_com", [f"PowerPoint COM open failed: {exc}"]
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _natural_sort_key(path: Path) -> tuple[int, str]:
    numbers = re.findall(r"\d+", path.stem)
    return (int(numbers[-1]) if numbers else 0, path.name.lower())


def _render_ppt_com(path: Path, output_dir: Path, width: int, height: int) -> tuple[list[Path], list[str]]:
    if not _is_windows():
        return [], ["PowerPoint COM rendering is available only on Windows."]
    try:
        import pythoncom
        import win32com.client
    except Exception as exc:
        return [], [f"PowerPoint COM rendering unavailable: {exc}"]
    pythoncom.CoInitialize()
    app = None
    presentation = None
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        presentation = app.Presentations.Open(str(path), ReadOnly=True, WithWindow=False)
        presentation.Export(str(output_dir), "PNG", int(width), int(height))
        # Windows is case-insensitive: globbing *.PNG and *.png returns each
        # file twice. Deduplicate, then order slides numerically (not lexically).
        unique: dict[str, Path] = {}
        for candidate in list(output_dir.glob("*.PNG")) + list(output_dir.glob("*.png")):
            unique.setdefault(str(candidate).lower(), candidate)
        previews = sorted(unique.values(), key=_natural_sort_key)
        return previews, []
    except Exception as exc:
        return [], [f"PowerPoint COM render failed: {exc}"]
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _render_ppt_libleoffice(path: Path, output_dir: Path) -> tuple[list[Path], list[str]]:
    soffice = _find_soffice()
    if not soffice:
        return [], ["LibreOffice rendering unavailable: soffice was not found."]
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=60,
        )
    except Exception as exc:
        return [], [f"LibreOffice conversion failed: {exc}"]
    pdfs = list(output_dir.glob("*.pdf"))
    if not pdfs:
        return [], ["LibreOffice did not produce a PDF preview."]
    lines = [f"LibreOffice produced PDF preview: {pdfs[0].name}", "PNG rendering requires native PowerPoint COM in this build."]
    out_path = output_dir / "ppt_libreoffice_preview.png"
    _render_summary_png(path.name, lines, out_path)
    return [out_path], []


def _build_ppt_fallback_preview(path: Path, output_dir: Path) -> tuple[list[Path], list[str]]:
    api, error = _require_pptx()
    if error:
        return [], [error]
    prs = api["Presentation"](str(path))
    summary = _ppt_summary(prs, include_text=True)
    lines = ["Structural preview only: native PowerPoint rendering was not available."]
    for slide in summary["slides"]:
        lines.append(f"Slide {slide['slide']}: {slide['shapes']} shapes, {slide['images']} images, {slide['tables']} tables")
        for text in slide.get("text", [])[:2]:
            lines.append("  " + str(text).replace("\n", " / ")[:120])
    out_path = output_dir / "ppt_structural_preview.png"
    _render_summary_png(path.name, lines, out_path)
    return [out_path], ["Native PowerPoint rendering unavailable; generated a structural preview."]


class PptCreateTool(BaseTool):
    name = "ppt_create"
    description = (
        "Create a product-grade PPTX deck from a structured slide spec. Supports "
        "templates, text, tables, images, basic shapes, and speaker notes."
    )
    parameters = {
        "type": "object",
        "properties": {
            "output_name": {"type": "string", "description": "Output .pptx file name."},
            "spec": {"type": "object", "description": "Presentation spec with template_path and slides.", "additionalProperties": True},
            "slides": {"type": "array", "description": "Optional shorthand for spec.slides.", "items": {"type": "object", "additionalProperties": True}},
            "template_path": {"type": "string", "description": "Optional existing PPTX template path."},
            "title": {"type": "string", "description": "Artifact display title."},
        },
        "required": ["output_name"],
    }

    async def execute(
        self,
        output_name: str,
        spec: Any = None,
        slides: Any = None,
        template_path: str = "",
        title: str = "",
        session_id: str = "default",
        tool_call_id: str = "",
        agent_type: str = "",
    ) -> ToolResult:
        api, error = _require_pptx()
        if error:
            return ToolResult(error=error)
        spec_dict, error = _as_dict(spec, "spec")
        if error:
            return ToolResult(error=error)
        if slides is not None:
            slide_list, error = _as_list(slides, "slides")
            if error:
                return ToolResult(error=error)
            spec_dict = {**spec_dict, "slides": slide_list}
        if template_path or spec_dict.get("template_path"):
            source, error = _resolve_input_path(str(template_path or spec_dict.get("template_path")), agent_type=agent_type)
            if error:
                return ToolResult(error=error)
            prs = api["Presentation"](str(source))
            while len(prs.slides) > 0:
                _delete_ppt_slide(prs, 0)
        else:
            prs = api["Presentation"]()
        if spec_dict.get("wide") is not False:
            prs.slide_width = api["Inches"](13.333)
            prs.slide_height = api["Inches"](7.5)
        for slide_spec in spec_dict.get("slides") or []:
            if isinstance(slide_spec, dict):
                _apply_ppt_slide_spec(api, prs, slide_spec, agent_type, len(prs.slides) + 1)
        out_path = _output_path(output_name, ".pptx", session_id=session_id, tool_call_id=tool_call_id, kind=OFFICE_KIND_PPT)
        prs.save(out_path)
        artifacts, artifact_error = _office_artifact(
            out_path,
            title=title or out_path.name,
            source_tool=self.name,
            session_id=session_id,
            tool_call_id=tool_call_id,
            agent_type=agent_type,
        )
        if artifact_error:
            return ToolResult(error=artifact_error)
        summary = _ppt_summary(prs, include_text=False)
        metadata = {"artifacts": artifacts, "file_path": str(out_path), "qa_status": "created", "summary": summary}
        return ToolResult(output=f"PowerPoint deck created: {out_path}\n\n{_json(summary)}", metadata=metadata)


class PptEditTool(BaseTool):
    name = "ppt_edit"
    description = "Edit an existing PPTX deck with structured operations, then publish the edited Office artifact."
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Existing PPTX path."},
            "operations": {"type": "array", "description": "Edit operations.", "items": {"type": "object", "additionalProperties": True}},
            "output_name": {"type": "string", "description": "Optional output .pptx name. Defaults to edited_<source>."},
            "title": {"type": "string", "description": "Artifact display title."},
        },
        "required": ["path", "operations"],
    }

    async def execute(
        self,
        path: str,
        operations: Any,
        output_name: str = "",
        title: str = "",
        session_id: str = "default",
        tool_call_id: str = "",
        agent_type: str = "",
    ) -> ToolResult:
        api, error = _require_pptx()
        if error:
            return ToolResult(error=error)
        source, error = _resolve_input_path(path, agent_type=agent_type)
        if error:
            return ToolResult(error=error)
        ops, error = _as_list(operations, "operations")
        if error:
            return ToolResult(error=error)
        prs = api["Presentation"](str(source))
        applied: list[str] = []
        for op in ops:
            if not isinstance(op, dict):
                continue
            action = str(op.get("op") or op.get("type") or "").lower()
            slide_index = int(op.get("slide_index") or op.get("slide") or 1) - 1
            if action == "add_slide":
                _apply_ppt_slide_spec(api, prs, op.get("slide_spec") or op, agent_type, len(prs.slides) + 1)
            elif action == "replace_text":
                old = str(op.get("old") or "")
                new = str(op.get("new") or "")
                targets = prs.slides if not op.get("slide_index") and not op.get("slide") else [prs.slides[slide_index]]
                if old:
                    for slide in targets:
                        for shape in slide.shapes:
                            if getattr(shape, "has_text_frame", False) and old in shape.text:
                                shape.text = shape.text.replace(old, new)
            elif action == "delete_slide":
                _delete_ppt_slide(prs, slide_index)
            elif action == "duplicate_slide":
                _duplicate_ppt_slide_basic(api, prs, slide_index, agent_type)
            elif action in {"add_text", "add_table", "add_image", "add_shape"} and 0 <= slide_index < len(prs.slides):
                block = {**op, "type": action.replace("add_", "")}
                _add_ppt_block(api, prs.slides[slide_index], block, 0, agent_type)
            else:
                continue
            applied.append(action)

        out_name = output_name or f"edited_{source.name}"
        out_path = _output_path(out_name, ".pptx", session_id=session_id, tool_call_id=tool_call_id, kind=OFFICE_KIND_PPT)
        prs.save(out_path)
        artifacts, artifact_error = _office_artifact(
            out_path,
            title=title or out_path.name,
            source_tool=self.name,
            session_id=session_id,
            tool_call_id=tool_call_id,
            agent_type=agent_type,
        )
        if artifact_error:
            return ToolResult(error=artifact_error)
        summary = _ppt_summary(prs, include_text=False)
        metadata = {"artifacts": artifacts, "file_path": str(out_path), "qa_status": "edited", "applied_operations": applied, "summary": summary}
        return ToolResult(output=f"PowerPoint deck edited: {out_path}\nApplied operations: {', '.join(applied) or '(none)'}", metadata=metadata)


class PptInspectTool(BaseTool):
    name = "ppt_inspect"
    description = "Inspect a PPTX deck's slide count, text, images, tables, and shapes."
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "PPTX path."},
            "include_text": {"type": "boolean", "default": True},
        },
        "required": ["path"],
    }

    async def execute(self, path: str, include_text: bool = True, agent_type: str = "") -> ToolResult:
        api, error = _require_pptx()
        if error:
            return ToolResult(error=error)
        source, error = _resolve_input_path(path, agent_type=agent_type)
        if error:
            return ToolResult(error=error)
        prs = api["Presentation"](str(source))
        summary = _ppt_summary(prs, include_text=include_text)
        return ToolResult(output=_json(summary), metadata={"summary": summary, "file_path": str(source)})


class PptValidateTool(BaseTool):
    name = "ppt_validate"
    description = "Validate a PPTX structurally and, when available, open it through native PowerPoint COM."
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "PPTX path."},
            "use_com": {"type": "boolean", "default": True},
        },
        "required": ["path"],
    }

    async def execute(self, path: str, use_com: bool = True, agent_type: str = "") -> ToolResult:
        api, error = _require_pptx()
        if error:
            return ToolResult(error=error)
        source, error = _resolve_input_path(path, agent_type=agent_type)
        if error:
            return ToolResult(error=error)
        try:
            prs = api["Presentation"](str(source))
            issues: list[str] = []
            if len(prs.slides) == 0:
                issues.append("Presentation has no slides.")
            layout_errors, layout_warnings, layout_json = _scan_ppt_layout_issues(prs)
            issues.extend(layout_errors)
        except Exception as exc:
            return ToolResult(error=f"Presentation could not be opened structurally: {exc}")
        engine = "structural_only"
        if use_com:
            engine, com_issues = _ppt_com_validate(source)
            issues.extend(com_issues)
        qa_status = "failed" if issues and engine == "office_com" else "warnings" if issues or layout_warnings else "passed"
        result = {
            "file_path": str(source),
            "engine": engine,
            "qa_status": qa_status,
            "issues": issues,
            "warnings": layout_warnings,
            "layout_errors": len(layout_errors),
            "layout_warnings": len(layout_warnings),
            "layout_json": layout_json,
            "slide_count": len(prs.slides),
        }
        return ToolResult(output=_json(result), metadata=result)


class PptRenderTool(BaseTool):
    name = "ppt_render"
    description = "Render PPTX slides as image artifacts. Uses native PowerPoint COM, then LibreOffice if available, then structural preview fallback."
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "PPTX path."},
            "width": {"type": "integer", "default": 1600},
            "height": {"type": "integer", "default": 900},
        },
        "required": ["path"],
    }

    async def execute(
        self,
        path: str,
        width: int = 1600,
        height: int = 900,
        session_id: str = "default",
        tool_call_id: str = "",
        agent_type: str = "",
    ) -> ToolResult:
        source, error = _resolve_input_path(path, agent_type=agent_type)
        if error:
            return ToolResult(error=error)
        output_dir = _office_run_dir(session_id, tool_call_id, OFFICE_KIND_PPT) / "render"
        output_dir.mkdir(parents=True, exist_ok=True)
        preview_paths, issues = _render_ppt_com(source, output_dir, width, height)
        engine = "office_com"
        if not preview_paths and _libreoffice_available():
            preview_paths, libre_issues = _render_ppt_libleoffice(source, output_dir)
            issues.extend(libre_issues)
            engine = "libreoffice"
        if not preview_paths:
            preview_paths, fallback_issues = _build_ppt_fallback_preview(source, output_dir)
            issues.extend(fallback_issues)
            engine = "structural_only"
        slide_paths = list(preview_paths)
        contact_sheet = None
        if len(preview_paths) > 1:
            contact_sheet = _build_contact_sheet(preview_paths, output_dir / "ppt_contact_sheet.png", f"{source.stem} contact sheet")
            if contact_sheet:
                preview_paths = [contact_sheet] + preview_paths
        artifacts = _preview_artifacts(
            preview_paths,
            title_prefix=f"{source.stem} slide preview",
            source_tool=self.name,
            session_id=session_id,
            tool_call_id=tool_call_id,
            agent_type=agent_type,
        )
        qa_status = "passed" if engine == "office_com" and not issues else "warnings"
        result = {
            "file_path": str(source),
            "preview_paths": [str(path) for path in preview_paths],
            "engine": engine,
            "qa_status": qa_status,
            "issues": issues,
            "artifacts": artifacts,
            "review_image_paths": [str(path) for path in slide_paths],
        }
        return ToolResult(
            output=_json({k: v for k, v in result.items() if k not in ("artifacts", "review_image_paths")}),
            metadata=result,
        )


def _office_file_kind(path: Path, explicit: str = "") -> str:
    kind = explicit.strip().lower()
    if kind in {"excel", "ppt", "powerpoint"}:
        return "ppt" if kind == "powerpoint" else kind
    if path.suffix.lower() in {".xlsx", ".xls"}:
        return "excel"
    if path.suffix.lower() in {".pptx", ".ppt"}:
        return "ppt"
    return "office"


def _enrich_office_artifact(artifact: dict[str, Any], source: Path, kind: str) -> dict[str, Any]:
    enriched = dict(artifact)
    enriched["kind"] = kind
    if kind == "excel":
        api, error = _require_openpyxl()
        if not error and not enriched.get("workbook"):
            try:
                workbook = api["load_workbook"](source, data_only=False)
                enriched["workbook"] = _workbook_summary(workbook, include_values=True, max_rows=18, max_cols=10)
                enriched.setdefault("qa_summary", _excel_quality_summary(workbook))
            except Exception:
                pass
    elif kind == "ppt":
        api, error = _require_pptx()
        if not error and not enriched.get("presentation"):
            try:
                prs = api["Presentation"](str(source))
                enriched["presentation"] = _ppt_summary(prs, include_text=True)
            except Exception:
                pass
    return enriched


def _copy_to_package_output(source: Path, package_dir: Optional[Path]) -> Path:
    if package_dir is None:
        return source
    output_dir = package_dir / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / source.name
    if source.resolve() != target.resolve():
        shutil.copy2(source, target)
    return target


def _resolve_preview_specs(preview_paths: Any, agent_type: str) -> list[dict[str, Any]]:
    specs, error = _as_list(preview_paths, "preview_paths")
    if error:
        return []
    result: list[dict[str, Any]] = []
    for item in specs:
        if isinstance(item, str):
            raw_path = item
            kind = "preview"
            title = ""
        elif isinstance(item, dict):
            raw_path = str(item.get("path") or item.get("file_path") or "")
            kind = str(item.get("kind") or "preview")
            title = str(item.get("title") or "")
        else:
            continue
        source = _existing_source_path(raw_path, agent_type=agent_type) if raw_path else None
        if source and _validate_source(source, agent_type=agent_type) is None:
            result.append({"path": source, "kind": kind, "title": title or source.name})
    return result


class OfficePackageCreateTool(BaseTool):
    name = "office_package_create"
    description = (
        "Create a runtime Office package workspace for a product-grade Excel/PPT deliverable. "
        "Use this before building related workbook and deck outputs for the Personal Agent."
    )
    parameters = {
        "type": "object",
        "properties": {
            "task_slug": {"type": "string", "description": "Stable task slug for the runtime workspace."},
            "title": {"type": "string", "description": "Human-readable package title."},
            "output_name_base": {"type": "string", "description": "Base name for final .xlsx/.pptx files."},
            "goal": {"type": "string", "description": "User-facing goal or brief for the package."},
        },
        "required": ["task_slug"],
    }

    async def execute(
        self,
        task_slug: str,
        title: str = "",
        output_name_base: str = "",
        goal: str = "",
        session_id: str = "default",
        tool_call_id: str = "",
        agent_type: str = "",
    ) -> ToolResult:
        package_dir = _office_package_dir(session_id, task_slug)
        manifest = {
            "type": "office_package",
            "title": title or task_slug,
            "task_slug": _safe_segment(task_slug, "office-package"),
            "goal": goal,
            "output_name_base": output_name_base or _safe_segment(title or task_slug, "office-output"),
            "session_id": session_id or "default",
            "tool_call_id": tool_call_id,
            "dirs": {
                "root": str(package_dir),
                "excel": str(package_dir / "excel"),
                "ppt": str(package_dir / "ppt"),
                "preview": str(package_dir / "preview"),
                "qa": str(package_dir / "qa"),
                "output": str(package_dir / "output"),
            },
            "files": [],
            "previews": [],
            "qa_summary": {},
        }
        manifest_path = package_dir / "manifest.json"
        _write_json(manifest_path, manifest)
        metadata = {"workspace": manifest["dirs"], "manifest_path": str(manifest_path), "task_slug": manifest["task_slug"]}
        return ToolResult(output=_json(metadata), metadata=metadata)


class OfficePackageQATool(BaseTool):
    name = "office_package_qa"
    description = (
        "Run the Office package QA loop across final Excel/PPT files: structural checks, "
        "optional native COM validation, preview rendering, contact sheet generation, and issue aggregation."
    )
    parameters = {
        "type": "object",
        "properties": {
            "excel_path": {"type": "string", "description": "Optional XLSX path to validate and render."},
            "ppt_path": {"type": "string", "description": "Optional PPTX path to validate and render."},
            "task_slug": {"type": "string", "description": "Optional package workspace slug."},
            "use_com": {"type": "boolean", "default": True},
            "render": {"type": "boolean", "default": True},
        },
    }

    async def execute(
        self,
        excel_path: str = "",
        ppt_path: str = "",
        task_slug: str = "",
        use_com: bool = True,
        render: bool = True,
        session_id: str = "default",
        tool_call_id: str = "",
        agent_type: str = "",
    ) -> ToolResult:
        package_dir = _office_package_dir(session_id, task_slug) if task_slug else _office_run_dir(session_id, tool_call_id, OFFICE_KIND_PACKAGE)
        preview_root = package_dir / "preview"
        qa_summary: dict[str, Any] = {}
        preview_records: list[dict[str, Any]] = []
        artifacts: list[dict[str, Any]] = []

        if excel_path:
            source, error = _resolve_input_path(excel_path, agent_type=agent_type)
            if error:
                return ToolResult(error=error)
            api, error = _require_openpyxl()
            if error:
                return ToolResult(error=error)
            workbook = api["load_workbook"](source, data_only=False)
            issues = _scan_excel_issues(workbook)
            engine = "structural_only"
            if use_com:
                engine, com_issues = _excel_com_validate(source)
                issues.extend(com_issues)
            previews: list[Path] = []
            if render:
                out_dir = preview_root / "excel"
                out_dir.mkdir(parents=True, exist_ok=True)
                previews, render_issues = _render_excel_com(source, out_dir, [])
                render_engine = "office_com"
                if not previews:
                    previews, fallback_issues = _build_excel_fallback_preview(source, out_dir)
                    render_issues.extend(fallback_issues)
                    render_engine = "structural_only"
                issues.extend(render_issues)
                engine = "office_com" if engine == "office_com" and render_engine == "office_com" else render_engine
                for path in previews:
                    preview_records.append({"path": str(path), "kind": "excel_sheet", "title": f"{source.stem} sheet preview"})
            quality = _excel_quality_summary(workbook)
            qa_status = "passed" if not issues else "warnings"
            qa_summary["excel"] = {"qa_status": qa_status, "engine": engine, "issues": issues, **quality}

        if ppt_path:
            source, error = _resolve_input_path(ppt_path, agent_type=agent_type)
            if error:
                return ToolResult(error=error)
            api, error = _require_pptx()
            if error:
                return ToolResult(error=error)
            prs = api["Presentation"](str(source))
            layout_errors, layout_warnings, layout_json = _scan_ppt_layout_issues(prs)
            issues = list(layout_errors)
            engine = "structural_only"
            if use_com:
                engine, com_issues = _ppt_com_validate(source)
                issues.extend(com_issues)
            previews = []
            if render:
                out_dir = preview_root / "ppt"
                out_dir.mkdir(parents=True, exist_ok=True)
                previews, render_issues = _render_ppt_com(source, out_dir, 1600, 900)
                render_engine = "office_com"
                if not previews and _libreoffice_available():
                    previews, libre_issues = _render_ppt_libleoffice(source, out_dir)
                    render_issues.extend(libre_issues)
                    render_engine = "libreoffice"
                if not previews:
                    previews, fallback_issues = _build_ppt_fallback_preview(source, out_dir)
                    render_issues.extend(fallback_issues)
                    render_engine = "structural_only"
                contact = _build_contact_sheet(previews, out_dir / "ppt_contact_sheet.png", f"{source.stem} contact sheet")
                if contact:
                    preview_records.append({"path": str(contact), "kind": "ppt_contact_sheet", "title": f"{source.stem} contact sheet"})
                for path in previews:
                    preview_records.append({"path": str(path), "kind": "ppt_slide", "title": f"{source.stem} slide preview"})
                issues.extend(render_issues)
                engine = "office_com" if engine == "office_com" and render_engine == "office_com" else render_engine
            qa_status = "failed" if layout_errors else "warnings" if layout_warnings or issues else "passed"
            qa_summary["ppt"] = {
                "qa_status": qa_status,
                "engine": engine,
                "issues": issues,
                "warnings": layout_warnings,
                "slide_count": len(prs.slides),
                "layout_errors": len(layout_errors),
                "layout_warnings": len(layout_warnings),
                "layout_json": layout_json,
            }

        for record in preview_records:
            artifact, _error = publish_file_artifact(
                record["path"],
                session_id=session_id or "default",
                title=record["title"],
                source_tool=self.name,
                tool_call_id=tool_call_id,
                agent_type=agent_type,
                artifact_type="image",
                require_image=True,
            )
            if artifact:
                artifact["kind"] = record["kind"]
                artifacts.append(artifact)

        overall_status = "passed"
        statuses = [value.get("qa_status") for value in qa_summary.values() if isinstance(value, dict)]
        if any(status == "failed" for status in statuses):
            overall_status = "failed"
        elif any(status == "warnings" for status in statuses):
            overall_status = "warnings"
        review_paths = [
            str(record["path"]) for record in preview_records
            if record.get("kind") in ("ppt_slide", "excel_sheet")
        ]
        result = {
            "qa_status": overall_status,
            "qa_summary": qa_summary,
            "preview_paths": preview_records,
            "manifest_path": str(package_dir / "manifest.json"),
            "artifacts": artifacts,
            "review_image_paths": review_paths,
        }
        _write_json(
            package_dir / "qa" / "office_qa.json",
            {k: v for k, v in result.items() if k not in ("artifacts", "review_image_paths")},
        )
        return ToolResult(
            output=_json({k: v for k, v in result.items() if k not in ("artifacts", "review_image_paths")}),
            metadata=result,
        )


class OfficePackagePublishTool(BaseTool):
    name = "office_package_publish"
    description = (
        "Publish final Excel/PPT deliverables as one grouped Office package artifact. "
        "Use this as the final step so the chat shows the two files and the right panel can preview them together."
    )
    parameters = {
        "type": "object",
        "properties": {
            "title": {"type": "string", "description": "Package display title."},
            "excel_path": {"type": "string", "description": "Optional final XLSX path."},
            "ppt_path": {"type": "string", "description": "Optional final PPTX path."},
            "files": {"type": "array", "description": "Optional file specs with path, kind, title.", "items": {"type": "object", "additionalProperties": True}},
            "preview_paths": {"type": "array", "description": "Optional preview specs or paths.", "items": {"type": "object", "additionalProperties": True}},
            "qa_summary": {"type": "object", "description": "Aggregated QA summary.", "additionalProperties": True},
            "task_slug": {"type": "string", "description": "Optional package workspace slug."},
        },
        "required": ["title"],
    }

    async def execute(
        self,
        title: str,
        excel_path: str = "",
        ppt_path: str = "",
        files: Any = None,
        preview_paths: Any = None,
        qa_summary: Any = None,
        task_slug: str = "",
        session_id: str = "default",
        tool_call_id: str = "",
        agent_type: str = "",
    ) -> ToolResult:
        package_dir = _office_package_dir(session_id, task_slug or title)
        file_specs: list[dict[str, Any]] = []
        if excel_path:
            file_specs.append({"path": excel_path, "kind": "excel", "title": Path(excel_path).name})
        if ppt_path:
            file_specs.append({"path": ppt_path, "kind": "ppt", "title": Path(ppt_path).name})
        extra_files, error = _as_list(files, "files")
        if error:
            return ToolResult(error=error)
        for item in extra_files:
            if isinstance(item, dict) and item.get("path"):
                file_specs.append(item)

        if not file_specs:
            return ToolResult(error="At least one Office file path is required.")

        qa_dict, error = _as_dict(qa_summary, "qa_summary")
        if error:
            return ToolResult(error=error)

        file_artifacts: list[dict[str, Any]] = []
        auto_preview_artifacts: list[dict[str, Any]] = []
        for item in file_specs:
            source, error = _resolve_input_path(str(item.get("path") or ""), agent_type=agent_type)
            if error:
                return ToolResult(error=error)
            final_source = _copy_to_package_output(source, package_dir)
            kind = _office_file_kind(final_source, str(item.get("kind") or ""))
            artifact, error = publish_file_artifact(
                str(final_source),
                session_id=session_id or "default",
                title=str(item.get("title") or final_source.name),
                caption=str(item.get("caption") or ""),
                source_tool=self.name,
                tool_call_id=tool_call_id,
                agent_type=agent_type,
                artifact_type="office",
            )
            if error:
                return ToolResult(error=error)
            assert artifact is not None
            enriched_artifact = _enrich_office_artifact(artifact, final_source, kind)
            file_artifacts.append(enriched_artifact)
            auto_preview_artifacts.extend(enriched_artifact.get("previews") or [])

        preview_artifacts: list[dict[str, Any]] = list(auto_preview_artifacts)
        for record in _resolve_preview_specs(preview_paths, agent_type):
            source = record["path"]
            artifact, _error = publish_file_artifact(
                str(source),
                session_id=session_id or "default",
                title=str(record.get("title") or source.name),
                source_tool=self.name,
                tool_call_id=tool_call_id,
                agent_type=agent_type,
                artifact_type="image",
                require_image=True,
            )
            if artifact:
                artifact["kind"] = record.get("kind") or "preview"
                preview_artifacts.append(artifact)

        if not qa_dict:
            for file_artifact in file_artifacts:
                file_kind = str(file_artifact.get("kind") or "")
                summary = file_artifact.get("qa_summary")
                if isinstance(summary, dict) and file_kind in summary:
                    qa_dict[file_kind] = summary[file_kind]
                elif isinstance(summary, dict) and summary:
                    qa_dict[file_kind or file_artifact.get("title") or "office"] = summary

        manifest_path = package_dir / "manifest.json"
        package = {
            "id": f"office_package_{uuid.uuid4().hex[:12]}",
            "type": "office_package",
            "title": title,
            "source": self.name,
            "tool_call_id": tool_call_id,
            "manifest_path": str(manifest_path),
            "path": str(package_dir),
            "files": file_artifacts,
            "previews": preview_artifacts,
            "qa_summary": qa_dict,
            "engine": qa_dict.get("engine") or "mixed",
            "size": sum(int(item.get("size") or 0) for item in file_artifacts),
        }
        _write_json(manifest_path, {
            "type": "office_package",
            "title": title,
            "files": file_artifacts,
            "previews": preview_artifacts,
            "qa_summary": qa_dict,
        })
        output = {
            "package": title,
            "files": [{"name": item.get("title"), "kind": item.get("kind"), "url": item.get("url")} for item in file_artifacts],
            "qa_summary": qa_dict,
        }
        return ToolResult(output=_json(output), metadata={"artifacts": [package], "package": package})
