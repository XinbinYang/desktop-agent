from __future__ import annotations

from pathlib import Path
from typing import Any


MAX_EXTRACT_ROWS = 500
MAX_EXTRACT_COLS = 80


def _color_to_hex(color: Any) -> str:
    if color is None:
        return ""
    value = getattr(color, "rgb", None)
    if isinstance(value, str) and value:
        if len(value) == 8:
            return f"#{value[-6:]}"
        if len(value) == 6:
            return f"#{value}"
    indexed = getattr(color, "indexed", None)
    if indexed is not None:
        return str(indexed)
    return ""


def _cell_payload(cell: Any) -> dict[str, Any]:
    raw_value = cell.value
    formula = raw_value if isinstance(raw_value, str) and raw_value.startswith("=") else None
    value = None if formula else raw_value
    fill = getattr(getattr(cell, "fill", None), "fgColor", None)
    font = getattr(cell, "font", None)
    alignment = getattr(cell, "alignment", None)
    return {
        "address": cell.coordinate,
        "value": value,
        "formula": formula,
        "data_type": getattr(cell, "data_type", ""),
        "number_format": getattr(cell, "number_format", ""),
        "style": {
            "bold": bool(getattr(font, "bold", False)),
            "italic": bool(getattr(font, "italic", False)),
            "font_color": _color_to_hex(getattr(font, "color", None)),
            "fill": _color_to_hex(fill),
            "align": getattr(alignment, "horizontal", None),
            "wrap": bool(getattr(alignment, "wrap_text", False)),
        },
    }


def _chart_title(chart: Any) -> str:
    title = getattr(chart, "title", None)
    if title is None:
        return ""
    try:
        paragraphs = title.tx.rich.p
        parts: list[str] = []
        for paragraph in paragraphs:
            for run in paragraph.r:
                text = getattr(run, "t", None)
                if text:
                    parts.append(str(text))
        return "".join(parts)
    except Exception:
        return str(title)


def _chart_payload(chart: Any, index: int) -> dict[str, Any]:
    anchor = getattr(chart, "anchor", None)
    marker = ""
    try:
        marker = f"{anchor._from.col + 1}:{anchor._from.row + 1}"
    except Exception:
        marker = ""
    return {
        "index": index,
        "type": chart.__class__.__name__,
        "title": _chart_title(chart),
        "series_count": len(getattr(chart, "series", []) or []),
        "anchor": marker,
    }


def _table_payload(table: Any) -> dict[str, Any]:
    return {
        "name": getattr(table, "displayName", "") or getattr(table, "name", ""),
        "ref": getattr(table, "ref", ""),
    }


def extract_workbook(path: str | Path, *, max_rows: int = MAX_EXTRACT_ROWS, max_cols: int = MAX_EXTRACT_COLS) -> dict[str, Any]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    source = Path(path)
    workbook = load_workbook(source, data_only=False, read_only=False)
    sheets: list[dict[str, Any]] = []
    formula_count = 0
    table_count = 0
    chart_count = 0
    validation_count = 0

    for sheet_index, worksheet in enumerate(workbook.worksheets):
        row_limit = min(int(max_rows), int(worksheet.max_row or 0))
        col_limit = min(int(max_cols), int(worksheet.max_column or 0))
        rows: list[list[dict[str, Any]]] = []
        for row in worksheet.iter_rows(min_row=1, max_row=row_limit, min_col=1, max_col=col_limit):
            payload_row: list[dict[str, Any]] = []
            for cell in row:
                payload = _cell_payload(cell)
                if payload.get("formula"):
                    formula_count += 1
                payload_row.append(payload)
            rows.append(payload_row)

        tables = [_table_payload(table) for table in getattr(worksheet, "tables", {}).values()]
        charts = [_chart_payload(chart, idx + 1) for idx, chart in enumerate(getattr(worksheet, "_charts", []) or [])]
        validations = getattr(getattr(worksheet, "data_validations", None), "dataValidation", []) or []
        table_count += len(tables)
        chart_count += len(charts)
        validation_count += len(validations)

        columns = [
            {
                "index": idx,
                "letter": get_column_letter(idx),
                "width": getattr(worksheet.column_dimensions.get(get_column_letter(idx)), "width", None),
            }
            for idx in range(1, col_limit + 1)
        ]
        row_heights = {
            str(idx): worksheet.row_dimensions[idx].height
            for idx in range(1, row_limit + 1)
            if worksheet.row_dimensions[idx].height is not None
        }
        sheets.append({
            "id": f"sheet-{sheet_index + 1}",
            "name": worksheet.title,
            "index": sheet_index,
            "max_row": int(worksheet.max_row or 0),
            "max_column": int(worksheet.max_column or 0),
            "row_limit": row_limit,
            "col_limit": col_limit,
            "truncated": bool((worksheet.max_row or 0) > row_limit or (worksheet.max_column or 0) > col_limit),
            "freeze_panes": str(worksheet.freeze_panes or ""),
            "merged_ranges": [str(item) for item in worksheet.merged_cells.ranges],
            "tables": tables,
            "charts": charts,
            "validations": [
                {
                    "type": getattr(item, "type", ""),
                    "sqref": str(getattr(item, "sqref", "")),
                    "formula1": getattr(item, "formula1", None),
                }
                for item in validations
            ],
            "columns": columns,
            "row_heights": row_heights,
            "rows": rows,
        })

    return {
        "type": "workbook",
        "file_name": source.name,
        "sheet_count": len(sheets),
        "formula_count": formula_count,
        "table_count": table_count,
        "chart_count": chart_count,
        "validation_count": validation_count,
        "sheets": sheets,
    }


def scan_formula_errors(workbook_payload: dict[str, Any]) -> list[dict[str, Any]]:
    errors: list[dict[str, Any]] = []
    tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!")
    for sheet in workbook_payload.get("sheets") or []:
        for row in sheet.get("rows") or []:
            for cell in row:
                text = str(cell.get("value") if cell.get("value") is not None else cell.get("formula") or "")
                if any(token in text for token in tokens):
                    errors.append({
                        "sheet": sheet.get("name"),
                        "address": cell.get("address"),
                        "value": text,
                    })
    return errors
