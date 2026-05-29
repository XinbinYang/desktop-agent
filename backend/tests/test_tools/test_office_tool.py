import pytest

from app.tools import CODING_AGENT_TOOLS, PERSONAL_AGENT_TOOLS, TOOLS_BY_NAME
from app.tools.artifact_tool import OfficePublishTool
from app.tools.office_tool import (
    ExcelCreateTool,
    ExcelRenderTool,
    ExcelValidateTool,
    OfficePackagePublishTool,
    PptCreateTool,
    PptRenderTool,
    PptValidateTool,
)


@pytest.fixture(autouse=True)
def force_structural_manifest_previews(monkeypatch):
    import app.office_artifacts.renderers as renderers

    monkeypatch.setattr(
        renderers,
        "_render_excel_com",
        lambda path, output_dir: ([], ["native render unavailable"]),
    )
    monkeypatch.setattr(
        renderers,
        "_render_ppt_com",
        lambda path, output_dir: ([], ["native render unavailable"]),
    )
    monkeypatch.setattr(
        renderers,
        "_libreoffice_pdf_render",
        lambda path, output_dir, prefix: ([], ["LibreOffice unavailable"]),
    )


@pytest.mark.asyncio
async def test_excel_create_analysis_profile_builds_dashboard_structure():
    result = await ExcelCreateTool().execute(
        output_name="analysis.xlsx",
        quality_profile="analysis_dashboard",
        spec={
            "title": "Excel + PPT Skill Self-Test Dashboard",
            "headers": ["Product", "2026 Revenue", "YoY Growth"],
            "rows": [
                ["Core", {"type": "currency", "value": 27.2}, {"type": "percent", "value": 0.277}],
                ["Pro", {"type": "currency", "value": 20.6}, {"type": "percent", "value": 0.411}],
            ],
        },
        session_id="office-analysis",
        agent_type="personal",
    )

    assert result.error == ""
    from openpyxl import load_workbook

    workbook = load_workbook(result.metadata["file_path"], data_only=False)
    assert workbook.sheetnames == ["Dashboard", "Data", "Assumptions", "Checks"]
    assert workbook["Dashboard"]["A1"].value == "Excel + PPT Skill Self-Test Dashboard"
    assert workbook["Data"]["B2"].value == 27.2
    assert workbook["Data"]["C2"].number_format == "0.0%"
    assert len(workbook["Dashboard"]._charts) == 1


@pytest.mark.asyncio
async def test_excel_create_publishes_valid_workbook():
    result = await ExcelCreateTool().execute(
        output_name="sales_model.xlsx",
        spec={
            "sheets": [{
                "name": "Summary",
                "data": [
                    ["Region", "Revenue", "Cost", "Profit"],
                    ["North", 120, 70, None],
                    ["South", 90, 55, None],
                ],
                "formulas": [
                    {"cell": "D2", "formula": "=B2-C2"},
                    {"cell": "D3", "formula": "=B3-C3"},
                ],
                "tables": [{"range": "A1:D3", "name": "SalesTable"}],
                "formats": [
                    {"range": "A1:D1", "bold": True, "fill_color": "#1F4E78", "font_color": "#FFFFFF"},
                    {"range": "B2:D3", "number_format": "#,##0"},
                ],
                "charts": [{
                    "type": "bar",
                    "title": "Revenue",
                    "values": "B1:B3",
                    "categories": "A2:A3",
                    "position": "F2",
                }],
                "freeze_panes": "A2",
                "auto_filter": True,
            }],
        },
        session_id="office-xlsx",
        tool_call_id="call-xlsx",
        agent_type="personal",
    )

    assert result.error == ""
    path = result.metadata["file_path"]
    artifact = result.metadata["artifacts"][0]
    assert artifact["type"] == "office"
    assert artifact["title"] == "sales_model.xlsx"
    assert artifact["kind"] == "excel"
    assert artifact["manifest_url"].endswith(f"/{artifact['id']}/manifest")
    assert artifact["workbook"]["sheet_count"] == 1
    assert artifact["previews"][0]["kind"] == "excel_sheet"

    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Summary"]
    assert sheet["A1"].value == "Region"
    assert sheet["D2"].value == "=B2-C2"
    assert "SalesTable" in sheet.tables
    assert len(sheet._charts) == 1


@pytest.mark.asyncio
async def test_excel_validate_structural_only_when_com_disabled():
    created = await ExcelCreateTool().execute(
        output_name="validate.xlsx",
        sheets=[{"name": "Data", "data": [["x"], [1]]}],
        session_id="office-validate",
        agent_type="personal",
    )

    result = await ExcelValidateTool().execute(
        path=created.metadata["file_path"],
        use_com=False,
        agent_type="personal",
    )

    assert result.error == ""
    assert result.metadata["qa_status"] == "passed"
    assert result.metadata["engine"] == "structural_only"


@pytest.mark.asyncio
async def test_excel_render_falls_back_to_structural_preview(monkeypatch):
    created = await ExcelCreateTool().execute(
        output_name="render.xlsx",
        sheets=[{"name": "Data", "data": [["x"], [1]]}],
        session_id="office-render",
        agent_type="personal",
    )

    import app.tools.office_tool as office_tool

    monkeypatch.setattr(
        office_tool,
        "_render_excel_com",
        lambda path, output_dir, sheet_names: ([], ["native render unavailable"]),
    )

    result = await ExcelRenderTool().execute(
        path=created.metadata["file_path"],
        session_id="office-render",
        agent_type="personal",
    )

    assert result.error == ""
    assert result.metadata["engine"] == "structural_only"
    artifact = result.metadata["artifacts"][0]
    assert artifact["type"] == "image"
    assert artifact["url"].startswith("/preview/office-render/")


@pytest.mark.asyncio
async def test_ppt_create_publishes_valid_deck():
    result = await PptCreateTool().execute(
        output_name="strategy_deck.pptx",
        spec={
            "slides": [
                {"layout": "title", "title": "Strategy Review", "subtitle": "Q2 update"},
                {
                    "layout": "blank",
                    "title": "Highlights",
                    "blocks": [
                        {"type": "bullet_list", "items": ["Revenue up", "Costs controlled"], "x": 0.9, "y": 1.5, "w": 8.5, "h": 1.3},
                        {"type": "table", "rows": [["Metric", "Value"], ["Revenue", "120"]], "x": 0.9, "y": 3.0, "w": 5.0, "h": 1.0},
                    ],
                },
            ],
        },
        session_id="office-ppt",
        tool_call_id="call-ppt",
        agent_type="personal",
    )

    assert result.error == ""
    path = result.metadata["file_path"]
    artifact = result.metadata["artifacts"][0]
    assert artifact["type"] == "office"
    assert artifact["title"] == "strategy_deck.pptx"
    assert artifact["kind"] == "ppt"
    assert artifact["presentation"]["slide_count"] == 2
    assert any(preview["kind"] == "ppt_slide" for preview in artifact["previews"])

    from pptx import Presentation

    deck = Presentation(path)
    assert len(deck.slides) == 2
    all_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    assert "Strategy Review" in all_text
    assert "Revenue up" in all_text


@pytest.mark.asyncio
async def test_ppt_validate_structural_only_when_com_disabled():
    created = await PptCreateTool().execute(
        output_name="validate_deck.pptx",
        slides=[{"title": "One slide"}],
        session_id="office-ppt-validate",
        agent_type="personal",
    )

    result = await PptValidateTool().execute(
        path=created.metadata["file_path"],
        use_com=False,
        agent_type="personal",
    )

    assert result.error == ""
    assert result.metadata["qa_status"] == "passed"
    assert result.metadata["engine"] == "structural_only"


@pytest.mark.asyncio
async def test_ppt_render_falls_back_to_structural_preview(monkeypatch):
    created = await PptCreateTool().execute(
        output_name="render_deck.pptx",
        slides=[{"title": "Preview slide"}],
        session_id="office-ppt-render",
        agent_type="personal",
    )

    import app.tools.office_tool as office_tool

    monkeypatch.setattr(
        office_tool,
        "_render_ppt_com",
        lambda path, output_dir, width, height: ([], ["native render unavailable"]),
    )
    monkeypatch.setattr(office_tool, "_libreoffice_available", lambda: False)

    result = await PptRenderTool().execute(
        path=created.metadata["file_path"],
        session_id="office-ppt-render",
        agent_type="personal",
    )

    assert result.error == ""
    assert result.metadata["engine"] == "structural_only"
    assert isinstance(result.metadata.get("review_image_paths"), list)
    artifact = result.metadata["artifacts"][0]
    assert artifact["type"] == "image"
    assert artifact["url"].startswith("/preview/office-ppt-render/")


@pytest.mark.asyncio
async def test_office_publish_accepts_xlsx_artifact():
    created = await ExcelCreateTool().execute(
        output_name="publish.xlsx",
        sheets=[{"name": "Data", "data": [["x"], [1]]}],
        session_id="office-publish",
        agent_type="personal",
    )

    result = await OfficePublishTool().execute(
        path=created.metadata["file_path"],
        title="Published workbook",
        session_id="office-publish",
        agent_type="personal",
    )

    assert result.error == ""
    artifact = result.metadata["artifacts"][0]
    assert artifact["type"] == "office"
    assert artifact["title"] == "Published workbook"


@pytest.mark.asyncio
async def test_office_package_publish_groups_excel_and_ppt_artifacts():
    workbook = await ExcelCreateTool().execute(
        output_name="package.xlsx",
        sheets=[{"name": "Data", "data": [["x"], [1]]}],
        session_id="office-package",
        agent_type="personal",
    )
    deck = await PptCreateTool().execute(
        output_name="package.pptx",
        slides=[{"title": "Package deck"}],
        session_id="office-package",
        agent_type="personal",
    )

    result = await OfficePackagePublishTool().execute(
        title="Office package",
        excel_path=workbook.metadata["file_path"],
        ppt_path=deck.metadata["file_path"],
        qa_summary={
            "excel": {"qa_status": "passed", "sheet_count": 1},
            "ppt": {"qa_status": "passed", "slide_count": 1},
        },
        task_slug="self-test",
        session_id="office-package",
        agent_type="personal",
    )

    assert result.error == ""
    package = result.metadata["artifacts"][0]
    assert package["type"] == "office_package"
    assert [file["kind"] for file in package["files"]] == ["excel", "ppt"]
    assert package["files"][0]["type"] == "office"
    assert package["files"][1]["type"] == "office"
    assert package["qa_summary"]["excel"]["qa_status"] == "passed"
    assert any(preview["kind"] == "excel_sheet" for preview in package["previews"])
    assert any(preview["kind"] == "ppt_slide" for preview in package["previews"])


@pytest.mark.asyncio
async def test_office_manifest_routes_and_xlsx_simple_edits(async_client):
    created = await ExcelCreateTool().execute(
        output_name="manifest_route.xlsx",
        spec={
            "sheets": [{
                "name": "Data",
                "data": [["Metric", "Value"], ["Revenue", 120]],
                "tables": [{"range": "A1:B2", "name": "MetricsTable"}],
            }],
        },
        session_id="office-route",
        agent_type="personal",
    )

    artifact = created.metadata["artifacts"][0]
    manifest_response = await async_client.get(f"/api/office/{artifact['id']}/manifest")
    assert manifest_response.status_code == 200
    manifest = manifest_response.json()
    assert manifest["version"] == 2
    assert manifest["sha256"]
    assert manifest["workbook"]["sheets"][0]["name"] == "Data"

    workbook_response = await async_client.get(f"/api/office/{artifact['id']}/workbook")
    assert workbook_response.status_code == 200
    assert workbook_response.json()["table_count"] == 1

    edit_response = await async_client.post(
        f"/api/office/{artifact['id']}/xlsx/simple-edits",
        json={"edits": [{"sheet": "Data", "cell": "B2", "value": 150}], "output_name": "edited.xlsx"},
    )
    assert edit_response.status_code == 200
    edited = edit_response.json()
    assert edited["type"] == "office"
    assert edited["kind"] == "excel"
    assert edited["workbook"]["sheets"][0]["rows"][1][1]["value"] == 150


def _elements_overlap(a: dict, b: dict) -> bool:
    dx = min(a["x"] + a["w"], b["x"] + b["w"]) - max(a["x"], b["x"])
    dy = min(a["y"] + a["h"], b["y"] + b["h"]) - max(a["y"], b["y"])
    return dx > 0.04 and dy > 0.04


def test_ppt_layout_engine_places_content_in_bounds_without_overlap():
    from app.office_artifacts.ppt_layout import SLIDE_H, SLIDE_W, plan_slides

    pages = plan_slides({
        "kicker": "组合",
        "title": "新秩序多空组合全景概述与战术配置说明",
        "subtitle": "做多 / 做空 / 对冲",
        "footer": "2026",
        "blocks": [
            {"type": "kpi_cards", "items": [
                {"value": "57%", "label": "做多"},
                {"value": "38%", "label": "做空"},
                {"value": "5%", "label": "尾部保护"},
            ]},
            {"type": "table", "header": True, "rows": [
                ["类别", "权重", "逻辑说明"],
                ["AI 基建", "20%", "芯片唯一代工与云业务高增长驱动"],
                ["久期/利率", "15%", "长久期利率敏感叠加商业地产承压"],
            ]},
            {"type": "bullets", "items": ["第一条核心要点说明", "第二条核心要点说明"]},
        ],
    })

    assert pages
    for page in pages:
        for element in page["elements"]:
            assert element["x"] >= -0.05 and element["y"] >= -0.05
            assert element["x"] + element["w"] <= SLIDE_W + 0.05
            assert element["y"] + element["h"] <= SLIDE_H + 0.05
        text_elements = [e for e in page["elements"] if e["kind"] == "text"]
        for i in range(len(text_elements)):
            for j in range(i + 1, len(text_elements)):
                assert not _elements_overlap(text_elements[i], text_elements[j])


def test_ppt_layout_paginates_when_content_overflows():
    from app.office_artifacts.ppt_layout import plan_slides

    blocks = [
        {"type": "bullets", "heading": f"段落 {i}",
         "items": ["要点甲" * 8, "要点乙" * 8, "要点丙" * 8]}
        for i in range(8)
    ]
    pages = plan_slides({"title": "长内容自动分页", "blocks": blocks})
    assert len(pages) >= 2


def test_ppt_layout_table_columns_fill_content_width():
    from app.office_artifacts.ppt_layout import CONTENT_W, DEFAULT_THEME, _prepare_table

    prepared = _prepare_table(
        {"type": "table", "header": True, "rows": [["A", "B", "C"], ["1", "2", "3"]]},
        CONTENT_W,
        DEFAULT_THEME,
    )
    table_element = prepared["emit"](0.72, 1.5, CONTENT_W)[0]
    assert abs(sum(table_element["col_w"]) - CONTENT_W) < 0.05


@pytest.mark.asyncio
async def test_ppt_create_semantic_spec_passes_layout_scan():
    from app.tools.office_tool import _require_pptx, _scan_ppt_layout_issues

    result = await PptCreateTool().execute(
        output_name="semantic_deck.pptx",
        spec={"slides": [
            {"layout": "cover", "kicker": "K", "title": "封面标题", "subtitle": "副标题"},
            {"kicker": "全景", "title": "组合全景", "subtitle": "概述",
             "blocks": [
                 {"type": "kpi_cards", "items": [
                     {"value": "57%", "label": "做多"},
                     {"value": "38%", "label": "做空"},
                 ]},
                 {"type": "table", "header": True, "rows": [
                     ["类别", "权重", "逻辑说明"],
                     ["AI 基建", "20%", "芯片唯一代工与云业务高增长驱动"],
                 ]},
             ]},
        ]},
        session_id="office-semantic",
        tool_call_id="call-semantic",
        agent_type="personal",
    )

    assert result.error == ""
    api, _ = _require_pptx()
    deck = api["Presentation"](result.metadata["file_path"])
    errors, _warnings, _layout = _scan_ppt_layout_issues(deck)
    assert errors == []


def test_office_tools_registered_for_personal_and_coding_agents():
    expected = {
        "office_publish",
        "office_package_create",
        "office_package_qa",
        "office_package_publish",
        "excel_inspect",
        "excel_create",
        "excel_edit",
        "excel_validate",
        "excel_render",
        "ppt_inspect",
        "ppt_create",
        "ppt_edit",
        "ppt_validate",
        "ppt_render",
    }

    assert expected <= set(TOOLS_BY_NAME)
    assert expected <= PERSONAL_AGENT_TOOLS
    assert expected <= CODING_AGENT_TOOLS
