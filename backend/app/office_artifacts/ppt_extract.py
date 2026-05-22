from __future__ import annotations

from pathlib import Path
from typing import Any


def _emu_to_px(value: Any) -> float:
    try:
        return round(float(value or 0) / 914400 * 96, 2)
    except Exception:
        return 0.0


def _emu_to_in(value: Any) -> float:
    try:
        return round(float(value or 0) / 914400, 4)
    except Exception:
        return 0.0


def _shape_bounds(shape: Any) -> dict[str, float]:
    return {
        "x": _emu_to_in(getattr(shape, "left", 0)),
        "y": _emu_to_in(getattr(shape, "top", 0)),
        "w": _emu_to_in(getattr(shape, "width", 0)),
        "h": _emu_to_in(getattr(shape, "height", 0)),
        "px_x": _emu_to_px(getattr(shape, "left", 0)),
        "px_y": _emu_to_px(getattr(shape, "top", 0)),
        "px_w": _emu_to_px(getattr(shape, "width", 0)),
        "px_h": _emu_to_px(getattr(shape, "height", 0)),
    }


def _table_rows(shape: Any) -> list[list[str]]:
    if not getattr(shape, "has_table", False):
        return []
    rows: list[list[str]] = []
    try:
        for row in shape.table.rows:
            rows.append([cell.text for cell in row.cells])
    except Exception:
        return []
    return rows


def _notes_text(slide: Any) -> str:
    try:
        return str(slide.notes_slide.notes_text_frame.text or "").strip()
    except Exception:
        return ""


def _shape_payload(shape: Any, index: int) -> dict[str, Any]:
    text = ""
    if getattr(shape, "has_text_frame", False):
        text = str(getattr(shape, "text", "") or "")
    return {
        "index": index,
        "name": str(getattr(shape, "name", "") or ""),
        "shape_type": str(getattr(shape, "shape_type", "") or ""),
        "bounds": _shape_bounds(shape),
        "text": text,
        "has_text": bool(text.strip()),
        "has_table": bool(getattr(shape, "has_table", False)),
        "has_image": bool(hasattr(shape, "image")),
        "table": _table_rows(shape),
    }


def _rect_overlap(a: dict[str, float], b: dict[str, float]) -> float:
    ax, ay, aw, ah = a["x"], a["y"], a["w"], a["h"]
    bx, by, bw, bh = b["x"], b["y"], b["w"], b["h"]
    dx = min(ax + aw, bx + bw) - max(ax, bx)
    dy = min(ay + ah, by + bh) - max(ay, by)
    if dx <= 0 or dy <= 0:
        return 0.0
    return dx * dy


def _layout_issues(slides: list[dict[str, Any]], slide_w: float, slide_h: float) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    for slide in slides:
        text_shapes = [shape for shape in slide["shapes"] if shape.get("text", "").strip()]
        if not slide["shapes"]:
            errors.append(f"Slide {slide['slide']} is empty.")
        for shape in slide["shapes"]:
            bounds = shape["bounds"]
            if bounds["x"] < -0.05 or bounds["y"] < -0.05 or bounds["x"] + bounds["w"] > slide_w + 0.05 or bounds["y"] + bounds["h"] > slide_h + 0.05:
                errors.append(f"Slide {slide['slide']} shape {shape['index']} is outside the slide bounds.")
        for left_index, left in enumerate(text_shapes):
            for right in text_shapes[left_index + 1:]:
                if left.get("has_table") or right.get("has_table"):
                    continue
                overlap = _rect_overlap(left["bounds"], right["bounds"])
                if overlap <= 0:
                    continue
                left_area = max(0.01, left["bounds"]["w"] * left["bounds"]["h"])
                right_area = max(0.01, right["bounds"]["w"] * right["bounds"]["h"])
                ratio = overlap / min(left_area, right_area)
                if ratio > 0.20 and overlap > 0.08:
                    errors.append(f"Slide {slide['slide']} text shapes {left['index']} and {right['index']} overlap badly.")
                elif ratio > 0.07 and overlap > 0.04:
                    warnings.append(f"Slide {slide['slide']} text shapes {left['index']} and {right['index']} overlap.")
    return errors, warnings


def extract_presentation(path: str | Path) -> dict[str, Any]:
    from pptx import Presentation

    source = Path(path)
    prs = Presentation(str(source))
    slide_w = _emu_to_in(prs.slide_width)
    slide_h = _emu_to_in(prs.slide_height)
    slides: list[dict[str, Any]] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        shapes = [_shape_payload(shape, idx + 1) for idx, shape in enumerate(slide.shapes)]
        texts = [shape["text"] for shape in shapes if shape.get("text", "").strip()]
        slides.append({
            "id": f"slide-{slide_index}",
            "slide": slide_index,
            "title": texts[0].splitlines()[0][:160] if texts else f"Slide {slide_index}",
            "texts": texts,
            "notes": _notes_text(slide),
            "shape_count": len(shapes),
            "image_count": sum(1 for shape in shapes if shape["has_image"]),
            "table_count": sum(1 for shape in shapes if shape["has_table"]),
            "shapes": shapes,
        })
    layout_errors, layout_warnings = _layout_issues(slides, slide_w, slide_h)
    return {
        "type": "presentation",
        "file_name": source.name,
        "slide_count": len(slides),
        "slide_width": slide_w,
        "slide_height": slide_h,
        "slides": slides,
        "layout_errors": layout_errors,
        "layout_warnings": layout_warnings,
    }
